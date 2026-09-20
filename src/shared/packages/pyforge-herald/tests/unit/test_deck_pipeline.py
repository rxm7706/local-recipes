"""``deck_pipeline.seed`` -- CAP-1 (Story 1.6): the local-prove gate, the
state-based conflict check, the 8-step write sequence, and the dual
state.py/registry.py record.

Every transport call is against a hand-written ``FakeTransport`` (no
network, no adapter); every local-prove call is against a hand-written
``FakeProver`` (no real ``npm`` subprocess).
"""

from __future__ import annotations

import base64
import hashlib
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import pytest
from pptx import Presentation

from pyforge.herald import deck_pipeline as deck_pipeline_module
from pyforge.herald import registry, stamps, state
from pyforge.herald.deck_pipeline import (
    PILOT_SUPPORT_SOURCE_PROJECT_ID,
    PROTOTYPE_ARTIFACT_KEY,
    STANDALONE_BUNDLE_ARTIFACT_KEY,
    AdoptedArtifact,
    ExportPushResult,
    NpmLocalProver,
    PixiDeckExporter,
    PptxTemplateExporter,
    PullResult,
    SeedResult,
    SubprocessGitCommitter,
    _persona_from_slug,
    _PixiPartialDeckExporter,
    _windowed_read,
    adopt,
    pull_marp_source,
    pull_prototype,
    pull_standalone_bundle,
    push_exports,
    seed,
    select_exporter,
)
from pyforge.herald.errors import (
    AuthError,
    ExportConflictError,
    HeraldError,
    PaginationStalledError,
    PptxTemplateError,
    ReadBackMismatchError,
    SeedConflictError,
    TransportCallError,
)
from pyforge.herald.registry import read as read_registry
from pyforge.herald.registry import register
from pyforge.herald.transport.base import FileRead, PlanHandle, PreviewRef, ProjectRef


class FakeTransport:
    """A hand-written ``DesignTransport`` double recording every call, in
    order, as ``(method, kwargs)``."""

    def __init__(self, *, prompt="PROMPT", project=None, plan=None, fails: dict | None = None):
        self.calls: list[tuple[str, dict]] = []
        self._prompt = prompt
        self._project = project or ProjectRef(project_id="p-new", url="https://claude.ai/design/p/p-new")
        self._plan = plan or PlanHandle(plan_token="tok", base_etags={"support.js": "0", "deck-stage.js": "0"})
        # Keyed by method name -- raises that exception INSTEAD OF the
        # normal canned return, after still recording the call.
        self._fails: dict = dict(fails or {})

    def get_design_prompt(self, **kwargs):
        self.calls.append(("get_design_prompt", kwargs))
        if "get_design_prompt" in self._fails:
            raise self._fails["get_design_prompt"]
        return self._prompt

    def create_project(self, **kwargs):
        self.calls.append(("create_project", kwargs))
        if "create_project" in self._fails:
            raise self._fails["create_project"]
        return self._project

    def finalize_plan(self, **kwargs):
        self.calls.append(("finalize_plan", kwargs))
        if "finalize_plan" in self._fails:
            raise self._fails["finalize_plan"]
        return self._plan

    def create_support_js(self, **kwargs):
        self.calls.append(("create_support_js", kwargs))
        return {}

    def copy_files(self, **kwargs):
        self.calls.append(("copy_files", kwargs))
        return {}

    def write_files(self, **kwargs):
        self.calls.append(("write_files", kwargs))
        return {}

    def read_file(self, **kwargs) -> FileRead:
        raise NotImplementedError("seed never reads")

    def render_preview(self, **kwargs) -> PreviewRef:
        raise NotImplementedError("seed never previews")

    def names(self) -> list[str]:
        return [name for name, _kwargs in self.calls]


class FakeProver:
    def __init__(self, *, fails: HeraldError | None = None):
        self.calls: list[Path] = []
        self._fails = fails

    def prove(self, deck_dir: Path) -> None:
        self.calls.append(deck_dir)
        if self._fails is not None:
            raise self._fails


def _make_deck(tmp_path: Path, slug: str, persona: str | None = None) -> Path:
    persona = persona or _persona_from_slug(slug)
    deck_dir = tmp_path / "presentations" / slug
    project_dir = deck_dir / "project"
    project_dir.mkdir(parents=True)
    (project_dir / f"PyForge {persona}.dc.html").write_text("<html>proto</html>", encoding="utf-8")
    (deck_dir / "README.md").write_text(f"# {slug}\n\nBody.\n", encoding="utf-8")
    return deck_dir


# --- _persona_from_slug ------------------------------------------------------


@pytest.mark.parametrize(
    ("slug", "persona"),
    [
        ("pyforge-warden", "Warden"),
        ("pyforge-cli-anything-hub", "Cli Anything Hub"),
        ("standalone", "Standalone"),
    ],
)
def test_persona_from_slug(slug, persona):
    assert _persona_from_slug(slug) == persona


# --- the happy path ----------------------------------------------------------


def test_seed_happy_path_calls_every_step_in_order(tmp_path: Path):
    deck_dir = _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    prover = FakeProver()

    result = seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)

    assert isinstance(result, SeedResult)
    assert result.persona == "Warden"
    assert result.prototype_filename == "PyForge Warden.dc.html"
    assert result.project.project_id == "p-new"
    assert prover.calls == [deck_dir]
    assert transport.names() == [
        "get_design_prompt",
        "create_project",
        "finalize_plan",
        "create_support_js",
        "copy_files",
        "write_files",
    ]


def test_seed_create_project_uses_the_persona_naming_convention(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-doctor")
    transport = FakeTransport()
    seed(transport, slug="pyforge-doctor", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[1]
    assert kwargs["name"] == "PyForge Doctor deck"


def test_seed_finalize_plan_declares_all_three_write_paths(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[2]
    assert set(kwargs["writes"]) == {
        "support.js",
        "deck-stage.js",
        "PyForge Warden.dc.html",
    }


def test_seed_create_support_js_uses_the_returned_base_etag(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    plan = PlanHandle(plan_token="tok", base_etags={"support.js": "E9"})
    transport = FakeTransport(plan=plan)
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[3]
    assert kwargs["if_match"] == "E9"
    assert kwargs["plan_token"] == "tok"


def test_seed_create_support_js_falls_back_to_fresh_etag(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    plan = PlanHandle(plan_token="tok", base_etags={})
    transport = FakeTransport(plan=plan)
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[3]
    assert kwargs["if_match"] == "0"


def test_seed_copy_files_uses_the_default_pilot_source_project(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[4]
    entry = kwargs["files"][0]
    assert entry["src_project_id"] == PILOT_SUPPORT_SOURCE_PROJECT_ID
    assert entry["src_path"] == "deck-stage.js"


def test_seed_copy_files_honours_an_explicit_source_project_override(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    seed(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=FakeProver(),
        support_source_project_id="custom-project-id",
    )
    _name, kwargs = transport.calls[4]
    assert kwargs["files"][0]["src_project_id"] == "custom-project-id"


def test_seed_write_files_carries_the_exact_local_prototype_bytes(tmp_path: Path):
    deck_dir = _make_deck(tmp_path, "pyforge-warden")
    on_disk = (deck_dir / "project" / "PyForge Warden.dc.html").read_text(encoding="utf-8")
    transport = FakeTransport()
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    _name, kwargs = transport.calls[5]
    entry = kwargs["files"][0]
    assert entry["path"] == "PyForge Warden.dc.html"
    assert entry["data"] == on_disk


def test_seed_records_state(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded == state.DeckState(project_id="p-new", etags={}, last_pull=None)


def test_seed_registers_the_readme_design_project_section(tmp_path: Path):
    deck_dir = _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    registered = read_registry(deck_dir / "README.md")
    assert registered.project_id == "p-new"
    assert registered.project_name == "PyForge Warden deck"
    assert registered.file_url == "https://claude.ai/design/p/p-new"


def test_seed_records_state_immediately_so_a_mid_pipeline_failure_never_duplicates_the_project(
    tmp_path: Path,
):
    """Review finding (Blind Hunter): `state.write` used to happen only
    after EVERY transport call succeeded. A failure anywhere between
    `create_project` and the final `state.write`/`registry.register` left
    the already-created remote project untracked by state.py's own
    conflict gate -- a retry's `existing_state is not None` check passed
    cleanly and `create_project` ran AGAIN, producing a second, orphaned
    duplicate project for the same slug. `state.write` now happens
    immediately after `create_project` succeeds, so this exact retry
    refuses instead."""
    _make_deck(tmp_path, "pyforge-warden")
    failing_transport = FakeTransport(fails={"finalize_plan": RuntimeError("simulated: etag conflict")})
    with pytest.raises(RuntimeError, match="etag conflict"):
        seed(
            failing_transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=FakeProver(),
        )
    # The project WAS created (that call succeeded) and state.py already
    # knows about it, even though the overall seed() call raised.
    assert "create_project" in failing_transport.names()
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded == state.DeckState(project_id="p-new", etags={}, last_pull=None)

    # A retry must refuse -- never call create_project a second time.
    retry_transport = FakeTransport()
    with pytest.raises(SeedConflictError, match="already seeded"):
        seed(
            retry_transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=FakeProver(),
        )
    assert retry_transport.calls == []


def test_seed_refuses_a_missing_readme_before_any_transport_call(tmp_path: Path):
    """Review finding (Edge Case Hunter): `registry.register` (called at
    the very end of a successful run) refuses against a missing
    `readme_path` -- "this module never fabricates a whole README from
    nothing" -- but nothing checked for that up front. The old sequence let
    a real project get created and `state.write` succeed, then failed on
    `registry.register`, leaving a slug permanently marked seeded in
    state.py with no way to complete registration short of manual
    intervention. Now refuses before touching the local prover or any
    transport call at all."""
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    (deck_dir / "project").mkdir(parents=True)
    (deck_dir / "project" / "PyForge Warden.dc.html").write_text("<html>proto</html>", encoding="utf-8")
    # Deliberately NO README.md.
    transport = FakeTransport()
    prover = FakeProver()
    with pytest.raises(HeraldError, match="no README.md"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert prover.calls == []
    assert transport.calls == []
    assert state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden") is None


# --- refusals and their "writes nothing" guarantee ---------------------------


def test_seed_refuses_a_missing_deck_directory(tmp_path: Path):
    transport = FakeTransport()
    with pytest.raises(HeraldError, match="no deck directory"):
        seed(transport, slug="pyforge-nope", repo_root=tmp_path, prover=FakeProver())
    assert transport.calls == []


def test_seed_refuses_when_already_seeded_before_any_call(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    state_path = tmp_path / state.DEFAULT_STATE_PATH
    state.write(
        state_path,
        "pyforge-warden",
        state.DeckState(project_id="already-there", etags={}, last_pull=None),
    )
    transport = FakeTransport()
    prover = FakeProver()
    with pytest.raises(SeedConflictError, match="already seeded"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert transport.calls == []
    assert prover.calls == []


def test_seed_refuses_via_the_registry_bootstrap_fallback_when_no_state_entry(
    tmp_path: Path,
):
    """The four pilot decks were seeded by hand before state.py existed: no
    state entry, only a well-formed README section. Re-seeding one through
    this CLI must not silently create a second Design project."""
    deck_dir = _make_deck(tmp_path, "pyforge-warden")
    register(
        readme_path=deck_dir / "README.md",
        project_name="PyForge Warden deck",
        project_id="existing-project-id",
        file_url="https://claude.ai/design/p/existing-project-id",
    )
    transport = FakeTransport()
    prover = FakeProver()
    with pytest.raises(SeedConflictError, match="existing-project-id"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert transport.calls == []
    assert prover.calls == []


def test_seed_refuses_a_malformed_registry_section_rather_than_treating_it_as_unseeded(
    tmp_path: Path,
):
    """A README carrying the heading with a body that doesn't match the
    canonical two-line shape (DW-1-5-1: true of all 13 pre-existing decks)
    is corruption, not 'nothing registered yet' -- must still refuse."""
    deck_dir = _make_deck(tmp_path, "pyforge-warden")
    (deck_dir / "README.md").write_text(
        "# pyforge-warden\n\n"
        "## Design project (the bridge's far end)\n\n"
        'Prototype lives in Claude Design project **"PyForge Warden deck"** '
        "(`existing-project-id`):\n"
        "https://claude.ai/design/p/existing-project-id\n"
        'Pull it into this deck with the MCP bridge ("pull warden") -- see\n'
        "`docs/specs/presentation-deck.md` § *The MCP bridge*.\n",
        encoding="utf-8",
    )
    transport = FakeTransport()
    prover = FakeProver()
    with pytest.raises(SeedConflictError, match="malformed"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert transport.calls == []
    assert prover.calls == []


def test_seed_refuses_a_missing_local_prototype_before_proving(tmp_path: Path):
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    (deck_dir / "project").mkdir(parents=True)
    (deck_dir / "README.md").write_text("# pyforge-warden\n", encoding="utf-8")
    transport = FakeTransport()
    prover = FakeProver()
    with pytest.raises(HeraldError, match="no local prototype"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert prover.calls == []
    assert transport.calls == []


def test_seed_propagates_a_prove_failure_before_any_transport_call(tmp_path: Path):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport()
    prover = FakeProver(fails=HeraldError("npm run extract exited 1"))
    with pytest.raises(HeraldError, match="npm run extract"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=prover)
    assert transport.calls == []


def test_seed_refuses_an_empty_design_prompt_before_creating_a_project(
    tmp_path: Path,
):
    _make_deck(tmp_path, "pyforge-warden")
    transport = FakeTransport(prompt="")
    with pytest.raises(HeraldError, match="empty design-system prompt"):
        seed(transport, slug="pyforge-warden", repo_root=tmp_path, prover=FakeProver())
    assert transport.names() == ["get_design_prompt"]


# --- NpmLocalProver's own error-mapping, subprocess.run patched out --------


def test_npm_local_prover_runs_extract_then_build(monkeypatch, tmp_path: Path):
    calls: list[list[str]] = []

    class _Completed:
        returncode = 0
        stdout = ""
        stderr = ""

    def _run(args, **kwargs):
        calls.append(args)
        return _Completed()

    monkeypatch.setattr(subprocess, "run", _run)
    NpmLocalProver().prove(tmp_path)
    assert calls == [["npm", "run", "extract"], ["npm", "run", "build"]]


def test_npm_local_prover_maps_nonzero_exit_to_a_herald_error(monkeypatch, tmp_path: Path):
    class _Completed:
        returncode = 1
        stdout = ""
        stderr = "slide count mismatch"

    monkeypatch.setattr(subprocess, "run", lambda *a, **k: _Completed())
    with pytest.raises(HeraldError, match="slide count mismatch"):
        NpmLocalProver().prove(tmp_path)


def test_npm_local_prover_maps_a_launch_oserror_to_a_herald_error(monkeypatch, tmp_path: Path):
    def _raise(*args, **kwargs):
        raise FileNotFoundError("no such file: npm")

    monkeypatch.setattr(subprocess, "run", _raise)
    with pytest.raises(HeraldError, match="could not run"):
        NpmLocalProver().prove(tmp_path)


def test_npm_local_prover_maps_a_timeout_to_a_herald_error(monkeypatch, tmp_path: Path):
    def _raise(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="npm", timeout=1)

    monkeypatch.setattr(subprocess, "run", _raise)
    with pytest.raises(HeraldError, match="exceeded"):
        NpmLocalProver(timeout=1).prove(tmp_path)


# --- CAP-2: pull_prototype (Story 2.1) ---------------------------------------


class FakePullTransport:
    """A hand-written ``DesignTransport`` double exercising only
    ``read_file`` -- every other method raises, since ``pull_prototype``
    never calls them."""

    def __init__(self, *, answers=None):
        self.calls: list[dict] = []
        # Consumed one per call (a list), or a single canned FileRead reused
        # for every call.
        self._answers = answers

    def read_file(self, **kwargs) -> FileRead:
        self.calls.append(kwargs)
        answer = self._answers
        if isinstance(answer, list):
            assert answer, "FakePullTransport ran out of canned read_file answers"
            return answer.pop(0)
        return answer

    def get_design_prompt(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls get_design_prompt")

    def create_project(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls create_project")

    def finalize_plan(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls finalize_plan")

    def create_support_js(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls create_support_js")

    def copy_files(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls copy_files")

    def write_files(self, **kwargs):
        raise NotImplementedError("pull_prototype never calls write_files")

    def render_preview(self, **kwargs) -> PreviewRef:
        raise NotImplementedError("pull_prototype never calls render_preview")


class FakeExporter:
    def __init__(self, *, fails: HeraldError | None = None):
        self.calls: list[tuple[str, Path]] = []
        self._fails = fails

    def export(self, *, slug: str, repo_root: Path) -> None:
        self.calls.append((slug, repo_root))
        if self._fails is not None:
            raise self._fails


class FakeCommitter:
    def __init__(self, *, fails: HeraldError | None = None):
        self.calls: list[dict] = []
        self._fails = fails

    def commit(self, *, repo_root: Path, paths: list[Path], message: str) -> None:
        self.calls.append({"repo_root": repo_root, "paths": list(paths), "message": message})
        if self._fails is not None:
            raise self._fails


_FIXED_NOW = datetime(2026, 8, 7, 12, 0, 0, tzinfo=timezone.utc)


def _seed_state(tmp_path: Path, slug: str, *, project_id="p-1", etags=None) -> None:
    state.write(
        tmp_path / state.DEFAULT_STATE_PATH,
        slug,
        state.DeckState(project_id=project_id, etags=dict(etags or {}), last_pull=None),
    )


def test_pull_prototype_short_circuits_on_unchanged_and_skips_every_downstream_step(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden", etags={PROTOTYPE_ARTIFACT_KEY: "E1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="E1", body=None, unchanged=True))
    prover = FakeProver()
    exporter = FakeExporter()

    result = pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=prover,
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )

    assert result == PullResult(
        slug="pyforge-warden",
        artifact=PROTOTYPE_ARTIFACT_KEY,
        local_path=None,
        unchanged=True,
        etag="E1",
        committed=False,
    )
    # Nothing downstream of the etag check ran.
    assert prover.calls == []
    assert exporter.calls == []
    assert not (tmp_path / "presentations" / "pyforge-warden" / "project").exists()
    # state.py is untouched -- last_pull stays None, proving no re-write happened.
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.last_pull is None


def test_pull_prototype_if_none_match_uses_the_last_seen_etag(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden", etags={PROTOTYPE_ARTIFACT_KEY: "E1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="E1", body=None, unchanged=True))
    pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=FakeProver(),
        exporter=FakeExporter(),
        now=lambda: _FIXED_NOW,
    )
    assert transport.calls[0]["if_none_match"] == "E1"
    assert transport.calls[0]["project_id"] == "p-1"
    assert transport.calls[0]["path"] == "PyForge Warden.dc.html"


def test_pull_prototype_writes_the_decoded_body_and_re_derives(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(
        answers=FileRead(path="x", etag="E2", body="<html>edited & saved</html>", unchanged=False)
    )
    prover = FakeProver()
    exporter = FakeExporter()

    result = pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=prover,
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )

    local_path = tmp_path / "presentations" / "pyforge-warden" / "project" / "PyForge Warden.dc.html"
    assert result.local_path == local_path
    assert result.unchanged is False
    assert result.etag == "E2"
    assert local_path.read_text(encoding="utf-8") == "<html>edited & saved</html>"
    assert prover.calls == [tmp_path / "presentations" / "pyforge-warden"]
    assert exporter.calls == [("pyforge-warden", tmp_path)]
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[PROTOTYPE_ARTIFACT_KEY] == "E2"
    assert recorded.last_pull == _FIXED_NOW.isoformat()


def test_pull_prototype_body_is_not_re_decoded(tmp_path: Path):
    """`FileRead.body` is already entity-decoded by
    `transport.base.parse_read_response` -- `pull_prototype` must use it
    verbatim. A body containing a literal `&amp;` (e.g. the file's own prior
    content already had one) must survive unchanged, not be decoded a
    second time into `&`."""
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(
        answers=FileRead(path="x", etag="E3", body="already &amp; decoded once", unchanged=False)
    )
    pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=FakeProver(),
        exporter=FakeExporter(),
        now=lambda: _FIXED_NOW,
    )
    local_path = tmp_path / "presentations" / "pyforge-warden" / "project" / "PyForge Warden.dc.html"
    assert local_path.read_text(encoding="utf-8") == "already &amp; decoded once"


def test_pull_prototype_refuses_a_truncated_answer_before_writing(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(
        answers=FileRead(
            path="x",
            etag="E4",
            body="partial",
            unchanged=False,
            first_line=1,
            last_line=5,
            total_lines=20,
        )
    )
    prover = FakeProver()
    exporter = FakeExporter()
    with pytest.raises(HeraldError, match="truncated"):
        pull_prototype(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=prover,
            exporter=exporter,
            now=lambda: _FIXED_NOW,
        )
    assert prover.calls == []
    assert exporter.calls == []
    assert not (tmp_path / "presentations" / "pyforge-warden" / "project").exists()


def test_pull_prototype_refuses_a_changed_answer_with_no_body(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="E5", body=None, unchanged=False))
    with pytest.raises(HeraldError, match="no body"):
        pull_prototype(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=FakeProver(),
            exporter=FakeExporter(),
            now=lambda: _FIXED_NOW,
        )


def test_pull_prototype_refuses_when_not_seeded(tmp_path: Path):
    transport = FakePullTransport(answers=None)
    with pytest.raises(HeraldError, match="herald deck seed"):
        pull_prototype(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=FakeProver(),
            exporter=FakeExporter(),
        )
    assert transport.calls == []


def test_pull_prototype_propagates_an_export_failure_after_the_write_lands(
    tmp_path: Path,
):
    """Review finding: the etag used to be recorded BEFORE export ran, so a
    failed export became permanently unrecoverable via retry -- a rerun's
    `if_none_match` would already match the just-recorded etag, the server
    would answer `{unchanged: true}`, and the pull would silently report
    "unchanged" for a re-derivation that never actually completed. The file
    write still lands (so the local file reflects the pulled content), but
    the etag is now recorded only after export succeeds -- so a retry after
    fixing the transient export failure genuinely re-attempts, rather than
    short-circuiting."""
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="E6", body="<html>x</html>", unchanged=False))
    exporter = FakeExporter(fails=HeraldError("deck-export exited 1"))
    with pytest.raises(HeraldError, match="deck-export exited 1"):
        pull_prototype(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            prover=FakeProver(),
            exporter=exporter,
            now=lambda: _FIXED_NOW,
        )
    # The write lands even though export later fails...
    local_path = tmp_path / "presentations" / "pyforge-warden" / "project" / "PyForge Warden.dc.html"
    assert local_path.read_text(encoding="utf-8") == "<html>x</html>"
    # ...but the etag is NOT recorded, so a retry can still re-attempt
    # re-derivation instead of short-circuiting as "unchanged" forever.
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert PROTOTYPE_ARTIFACT_KEY not in recorded.etags


# --- CAP-2: --commit (Story 2.2) ---------------------------------------------


def test_pull_prototype_without_commit_never_calls_the_committer(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="E7", body="<html>x</html>", unchanged=False))
    committer = FakeCommitter()
    result = pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=FakeProver(),
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert committer.calls == []
    assert result.committed is False


def test_pull_prototype_commit_true_stages_and_commits_after_a_real_change(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="E8", body="<html>x</html>", unchanged=False))
    committer = FakeCommitter()
    result = pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        commit=True,
        prover=FakeProver(),
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert result.committed is True
    assert len(committer.calls) == 1
    call = committer.calls[0]
    assert call["repo_root"] == tmp_path
    assert tmp_path / "presentations" / "pyforge-warden" in call["paths"]
    assert tmp_path / state.DEFAULT_STATE_PATH in call["paths"]
    assert "pyforge-warden" in call["message"]


def test_pull_prototype_commit_true_never_commits_an_unchanged_pull(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden", etags={PROTOTYPE_ARTIFACT_KEY: "E1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="E1", body=None, unchanged=True))
    committer = FakeCommitter()
    result = pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        commit=True,
        prover=FakeProver(),
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert committer.calls == []
    assert result.committed is False
    assert result.unchanged is True


def test_pull_prototype_propagates_a_commit_failure(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="E9", body="<html>x</html>", unchanged=False))
    committer = FakeCommitter(fails=HeraldError("git commit failed: nothing to commit"))
    with pytest.raises(HeraldError, match="nothing to commit"):
        pull_prototype(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            commit=True,
            prover=FakeProver(),
            exporter=FakeExporter(),
            committer=committer,
            now=lambda: _FIXED_NOW,
        )
    # The write and state update already landed before the commit failed --
    # not rolled back (see the story spec's Design Notes).
    local_path = tmp_path / "presentations" / "pyforge-warden" / "project" / "PyForge Warden.dc.html"
    assert local_path.read_text(encoding="utf-8") == "<html>x</html>"


# --- CAP-2: pull_marp_source (Story 2.3) -------------------------------------


def test_pull_marp_source_short_circuits_on_unchanged(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden", etags={"marp:deck": "M1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="M1", body=None, unchanged=True))
    exporter = FakeExporter()
    result = pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="deck",
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )
    assert result == PullResult(
        slug="pyforge-warden",
        artifact="marp:deck",
        local_path=None,
        unchanged=True,
        etag="M1",
        committed=False,
    )
    assert exporter.calls == []
    assert not (tmp_path / "presentations" / "pyforge-warden" / "src").exists()


def test_pull_marp_source_uses_the_short_name_remote_path(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="M2", body="# Deck", unchanged=False))
    pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="executive-summary",
        exporter=FakeExporter(),
        now=lambda: _FIXED_NOW,
    )
    assert transport.calls[0]["path"] == "warden-executive-summary.md"


def test_pull_marp_source_lands_at_the_dated_src_marp_path_and_calls_no_prover(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="M3", body="# Infographic", unchanged=False))
    exporter = FakeExporter()

    result = pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="infographic",
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )

    local_path = (
        tmp_path / "presentations" / "pyforge-warden" / "src" / "marp" / "pyforge-warden-infographic-2026-08-07.md"
    )
    assert result.local_path == local_path
    assert result.artifact == "marp:infographic"
    assert local_path.read_text(encoding="utf-8") == "# Infographic"
    assert exporter.calls == [("pyforge-warden", tmp_path)]
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags["marp:infographic"] == "M3"


def test_pull_marp_source_refuses_an_unknown_kind(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=None)
    with pytest.raises(HeraldError, match="unknown Marp source kind"):
        pull_marp_source(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            kind="cover",
            exporter=FakeExporter(),
        )
    assert transport.calls == []


def test_pull_marp_source_refuses_when_not_seeded(tmp_path: Path):
    transport = FakePullTransport(answers=None)
    with pytest.raises(HeraldError, match="herald deck seed"):
        pull_marp_source(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            kind="deck",
            exporter=FakeExporter(),
        )
    assert transport.calls == []


def test_pull_marp_source_commit_true_stages_after_a_real_change(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="M4", body="# Deck", unchanged=False))
    committer = FakeCommitter()
    result = pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="deck",
        commit=True,
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert result.committed is True
    assert len(committer.calls) == 1
    assert "marp:deck" in committer.calls[0]["message"]


def test_pull_marp_source_commit_true_never_commits_an_unchanged_pull(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden", etags={"marp:deck": "M1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="M1", body=None, unchanged=True))
    committer = FakeCommitter()
    result = pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="deck",
        commit=True,
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert committer.calls == []
    assert result.committed is False


# --- CAP-2: pull_standalone_bundle (Story 2.4) -------------------------------


def test_pull_standalone_bundle_short_circuits_on_unchanged(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden", etags={STANDALONE_BUNDLE_ARTIFACT_KEY: "S1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="S1", body=None, unchanged=True))
    exporter = FakeExporter()
    result = pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )
    assert result == PullResult(
        slug="pyforge-warden",
        artifact=STANDALONE_BUNDLE_ARTIFACT_KEY,
        local_path=None,
        unchanged=True,
        etag="S1",
        committed=False,
    )
    assert exporter.calls == []
    assert not (tmp_path / "presentations" / "pyforge-warden" / "src").exists()


def test_pull_standalone_bundle_uses_the_persona_remote_path(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="S2", body="<html>bundle</html>", unchanged=False))
    pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        exporter=FakeExporter(),
        now=lambda: _FIXED_NOW,
    )
    assert transport.calls[0]["path"] == "Warden Infographic standalone.html"


def test_pull_standalone_bundle_lands_at_the_dated_export_path(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="S3", body="<html>bundle</html>", unchanged=False))
    exporter = FakeExporter()

    result = pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        exporter=exporter,
        now=lambda: _FIXED_NOW,
    )

    local_path = (
        tmp_path
        / "presentations"
        / "pyforge-warden"
        / "src"
        / "marp"
        / "pyforge-warden-infographic-standalone-2026-08-07.html"
    )
    assert result.local_path == local_path
    assert result.artifact == STANDALONE_BUNDLE_ARTIFACT_KEY
    assert local_path.read_text(encoding="utf-8") == "<html>bundle</html>"
    assert exporter.calls == [("pyforge-warden", tmp_path)]
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[STANDALONE_BUNDLE_ARTIFACT_KEY] == "S3"


def test_pull_standalone_bundle_write_completes_before_export_runs(tmp_path: Path):
    """The bundle-supersedes-marp-render logic belongs to `deck-export`
    (out of scope); this module's own responsibility is to guarantee no
    check-then-act race by finishing the atomic write strictly before
    invoking the exporter. Asserted here by an exporter fake that reads the
    file back at call time -- if the write had not completed, this would
    read stale/missing content instead of the just-pulled body."""
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="S4", body="<html>bundle</html>", unchanged=False))
    seen_at_export_time = {}

    class ReadingExporter:
        def export(self, *, slug: str, repo_root: Path) -> None:
            local_path = (
                repo_root
                / "presentations"
                / slug
                / "src"
                / "marp"
                / "pyforge-warden-infographic-standalone-2026-08-07.html"
            )
            seen_at_export_time["body"] = local_path.read_text(encoding="utf-8")

    pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        exporter=ReadingExporter(),
        now=lambda: _FIXED_NOW,
    )
    assert seen_at_export_time["body"] == "<html>bundle</html>"


def test_pull_standalone_bundle_refuses_when_not_seeded(tmp_path: Path):
    transport = FakePullTransport(answers=None)
    with pytest.raises(HeraldError, match="herald deck seed"):
        pull_standalone_bundle(
            transport,
            slug="pyforge-warden",
            repo_root=tmp_path,
            exporter=FakeExporter(),
        )
    assert transport.calls == []


def test_pull_standalone_bundle_commit_true_stages_after_a_real_change(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="S5", body="<html>bundle</html>", unchanged=False))
    committer = FakeCommitter()
    result = pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        commit=True,
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert result.committed is True
    assert len(committer.calls) == 1
    assert STANDALONE_BUNDLE_ARTIFACT_KEY in committer.calls[0]["message"]


def test_pull_standalone_bundle_commit_true_never_commits_an_unchanged_pull(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden", etags={STANDALONE_BUNDLE_ARTIFACT_KEY: "S1"})
    transport = FakePullTransport(answers=FileRead(path="x", etag="S1", body=None, unchanged=True))
    committer = FakeCommitter()
    result = pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        commit=True,
        exporter=FakeExporter(),
        committer=committer,
        now=lambda: _FIXED_NOW,
    )
    assert committer.calls == []
    assert result.committed is False


# --- SubprocessGitCommitter (real subprocess, real scratch git repo) --------


def _init_git_repo(work: Path) -> None:
    subprocess.run(["git", "init", "-q"], cwd=work, check=True)
    subprocess.run(["git", "config", "user.email", "test@example.com"], cwd=work, check=True)
    subprocess.run(["git", "config", "user.name", "Test"], cwd=work, check=True)
    (work / "README.md").write_text("scratch repo\n", encoding="utf-8")
    subprocess.run(["git", "add", "-A"], cwd=work, check=True)
    subprocess.run(["git", "commit", "-q", "-m", "init"], cwd=work, check=True)


def test_subprocess_git_committer_commits_with_an_absolute_repo_root(tmp_path: Path):
    work = tmp_path / "repo"
    work.mkdir()
    _init_git_repo(work)
    target = work / "presentations" / "warden"
    target.mkdir(parents=True)
    (target / "file.txt").write_text("content\n", encoding="utf-8")

    SubprocessGitCommitter().commit(repo_root=work, paths=[target], message="herald: pull warden (prototype)")

    log = subprocess.run(
        ["git", "log", "-1", "--name-only", "--format="],
        cwd=work,
        capture_output=True,
        text=True,
        check=True,
    )
    assert "presentations/warden/file.txt" in log.stdout


def test_subprocess_git_committer_commits_with_a_relative_repo_root(tmp_path: Path, monkeypatch):
    """Review finding: `p.is_absolute()` was the wrong branch condition --
    every `paths` entry callers pass is already prefixed with `repo_root`,
    whether or not `repo_root` itself is absolute. Invoking with a
    RELATIVE `repo_root` (e.g. `--repo-root some/subdir`) used to double
    the prefix (`some/subdir/some/subdir/...`), failing with a git
    pathspec error even though the pull itself succeeded."""
    work = tmp_path / "repo"
    work.mkdir()
    _init_git_repo(work)
    target = work / "presentations" / "warden"
    target.mkdir(parents=True)
    (target / "file.txt").write_text("content\n", encoding="utf-8")

    monkeypatch.chdir(tmp_path)
    relative_root = Path("repo")  # relative to the new cwd, resolves to `work`

    SubprocessGitCommitter().commit(
        repo_root=relative_root,
        paths=[relative_root / "presentations" / "warden"],
        message="herald: pull warden (prototype)",
    )

    log = subprocess.run(
        ["git", "log", "-1", "--name-only", "--format="],
        cwd=work,
        capture_output=True,
        text=True,
        check=True,
    )
    assert "presentations/warden/file.txt" in log.stdout


# --- CAP-5: push_exports (Story 5.1/5.2) -------------------------------------


class FakePushTransport:
    """A hand-written ``DesignTransport`` double exercising
    ``finalize_plan``/``write_files`` and (Story 23.4) ``fetch_rendered_bytes``
    -- every other method raises, since ``push_exports`` never calls them.

    ``write_fails`` maps a filename to the exception ``write_files`` should
    raise for that entry -- how Story 5.2's per-file conflict is simulated
    (the real wire shape for a conditional-write rejection is unproven, per
    DW-1-2-5; ``errors.TransportError`` is what a real rejection would
    surface through ``McpTransport``'s ``_call_json``/``require_conditional``
    failure path). ``rendered_bytes`` maps a filename to the bytes
    ``fetch_rendered_bytes`` answers with for ``--prove``'s read-back (a
    filename absent from the map answers ``b""``, deliberately a mismatch
    unless a test means for it to be); ``fetch_fails`` mirrors
    ``write_fails`` for that same call."""

    def __init__(
        self,
        *,
        plan: PlanHandle | None = None,
        write_fails: dict[str, Exception] | None = None,
        rendered_bytes: dict[str, bytes] | None = None,
        fetch_fails: dict[str, Exception] | None = None,
    ):
        self.finalize_plan_calls: list[dict] = []
        self.write_files_calls: list[dict] = []
        self.fetch_rendered_bytes_calls: list[dict] = []
        self._plan = plan or PlanHandle(plan_token="tok", base_etags={})
        self._write_fails = dict(write_fails or {})
        self._rendered_bytes = dict(rendered_bytes or {})
        self._fetch_fails = dict(fetch_fails or {})

    def finalize_plan(self, **kwargs):
        self.finalize_plan_calls.append(kwargs)
        return self._plan

    def write_files(self, **kwargs):
        self.write_files_calls.append(kwargs)
        path = kwargs["files"][0]["path"]
        if path in self._write_fails:
            raise self._write_fails[path]
        return {}

    def fetch_rendered_bytes(self, *, project_id: str, path: str) -> bytes:
        self.fetch_rendered_bytes_calls.append({"project_id": project_id, "path": path})
        if path in self._fetch_fails:
            raise self._fetch_fails[path]
        return self._rendered_bytes.get(path, b"")

    def get_design_prompt(self, **kwargs):
        raise NotImplementedError("push_exports never calls get_design_prompt")

    def create_project(self, **kwargs):
        raise NotImplementedError("push_exports never calls create_project")

    def create_support_js(self, **kwargs):
        raise NotImplementedError("push_exports never calls create_support_js")

    def copy_files(self, **kwargs):
        raise NotImplementedError("push_exports never calls copy_files")

    def read_file(self, **kwargs):
        raise NotImplementedError("push_exports never calls read_file")

    def render_preview(self, **kwargs):
        raise NotImplementedError("push_exports never calls render_preview")

    def list_files(self, **kwargs):
        raise NotImplementedError("push_exports never calls list_files")


def _write_export_html(tmp_path: Path, slug: str, date: str, body: str) -> Path:
    marp_dir = tmp_path / "presentations" / slug / "src" / "marp"
    marp_dir.mkdir(parents=True, exist_ok=True)
    path = marp_dir / f"{slug}-infographic-standalone-{date}.html"
    path.write_text(body, encoding="utf-8")
    return path


def _write_export_pptx(tmp_path: Path, slug: str, kind: str, date: str, raw_bytes: bytes) -> Path:
    """``kind`` is ``"deck"`` or ``"infographic"`` -- the two PPTX filename
    shapes ``scripts/deck_export.py`` produces (Story 23.4). The content is
    arbitrary bytes, never a real ``.pptx`` zip: nothing in
    ``_discover_export_files``/``push_exports`` parses PPTX structure, only
    reads and hashes raw bytes."""
    pptx_dir = tmp_path / "presentations" / slug / "src" / "pptx"
    pptx_dir.mkdir(parents=True, exist_ok=True)
    name = {
        "deck": f"{slug}-deck-{date}.pptx",
        "infographic": f"{slug}_infographic_deck-{date}.pptx",
    }[kind]
    path = pptx_dir / name
    path.write_bytes(raw_bytes)
    return path


def test_push_exports_first_push_uses_the_0_sentinel_etag(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    assert result == ExportPushResult(slug="pyforge-warden", pushed=(filename,), skipped=())
    assert transport.finalize_plan_calls == [{"project_id": "p-1", "writes": [filename]}]
    write_call = transport.write_files_calls[0]
    assert write_call["files"][0]["path"] == filename
    assert write_call["files"][0]["data"] == "<html>v1</html>"
    assert write_call["files"][0]["if_match"] == "0"


def test_push_exports_records_the_local_hash_after_a_successful_push(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    push_exports(FakePushTransport(), slug="pyforge-warden", repo_root=tmp_path)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    key = f"export:{filename}"
    assert key in recorded.etags
    assert recorded.etags[key] == hashlib.sha256(b"<html>v1</html>").hexdigest()


def test_push_exports_skips_a_file_whose_local_hash_is_unchanged(tmp_path: Path):
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    stored_hash = hashlib.sha256(b"<html>v1</html>").hexdigest()
    _seed_state(tmp_path, "pyforge-warden", etags={f"export:{filename}": stored_hash})
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result == ExportPushResult(slug="pyforge-warden", pushed=(), skipped=(filename,))
    assert transport.finalize_plan_calls == []
    assert transport.write_files_calls == []


def test_push_exports_pushes_a_changed_file_even_with_a_prior_record(tmp_path: Path):
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    _seed_state(tmp_path, "pyforge-warden", etags={f"export:{filename}": "stale-hash"})
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v2</html>")
    transport = FakePushTransport(plan=PlanHandle(plan_token="tok", base_etags={filename: "E9"}))

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result.pushed == (filename,)
    assert transport.write_files_calls[0]["files"][0]["if_match"] == "E9"
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    expected_hash = hashlib.sha256(b"<html>v2</html>").hexdigest()
    assert recorded.etags[f"export:{filename}"] == expected_hash


def test_push_exports_nothing_to_push_makes_no_transport_calls(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result == ExportPushResult(slug="pyforge-warden", pushed=(), skipped=())
    assert transport.finalize_plan_calls == []
    assert transport.write_files_calls == []


def test_push_exports_refuses_when_not_seeded(tmp_path: Path):
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport()

    with pytest.raises(HeraldError, match="cannot push .*herald deck seed"):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)
    assert transport.finalize_plan_calls == []


def test_discover_export_files_ignores_a_non_dated_stray_file(tmp_path: Path):
    """Regression: a plain lexicographic sort put a stray same-prefix file
    (a hand-copied backup, an aborted draft) AFTER every real dated file --
    letters sort after digits -- so `matches[-1]` silently picked the
    stale/unrelated file over the genuine current export. The date segment
    must actually parse as `YYYY-MM-DD` to be considered a candidate."""
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>REAL</html>")
    stray = (
        tmp_path
        / "presentations"
        / "pyforge-warden"
        / "src"
        / "marp"
        / "pyforge-warden-infographic-standalone-old-backup.html"
    )
    stray.write_text("<html>STALE</html>", encoding="utf-8")

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    assert len(candidates) == 1
    assert candidates[0].filename == "pyforge-warden-infographic-standalone-2026-08-07.html"
    assert candidates[0].data == "<html>REAL</html>"


def test_discover_export_files_picks_the_newest_of_several_dated_files(tmp_path: Path):
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-01", "<html>older</html>")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>newer</html>")

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    assert len(candidates) == 1
    assert candidates[0].data == "<html>newer</html>"


def test_push_exports_no_derived_file_on_disk_yet_is_a_no_op(tmp_path: Path):
    """`deck-export` has never run for this deck -- nothing to push, not an
    error (mirrors the pull-side "nothing changed" no-op ethos)."""
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result == ExportPushResult(slug="pyforge-warden", pushed=(), skipped=())


# --- CAP-5: push_exports conflict handling (Story 5.2) -----------------------


def test_push_exports_conflict_raises_export_conflict_error(tmp_path: Path):
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport(write_fails={filename: TransportCallError("write_files: etag mismatch")})

    with pytest.raises(ExportConflictError, match=filename):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)


def test_push_exports_conflict_does_not_record_state_for_the_conflicted_file(
    tmp_path: Path,
):
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport(write_fails={filename: TransportCallError("write_files: etag mismatch")})

    with pytest.raises(ExportConflictError):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert f"export:{filename}" not in recorded.etags


def test_push_exports_conflict_on_one_file_does_not_abort_the_rest_of_the_batch(tmp_path: Path, monkeypatch):
    """`_discover_export_files` only ever surfaces one candidate today (the
    scope note in `deck_pipeline.py`'s own CAP-5 section), so this test
    exercises the batch-continues-past-a-conflict loop directly by
    monkeypatching discovery to return two candidates -- proving the loop
    itself (not just today's single-file discovery) honors FR-20/NFR-02:
    a conflict on one file must not clobber, or block, another file's own
    successful push."""
    _seed_state(tmp_path, "pyforge-warden")
    bad_candidate = deck_pipeline_module._ExportCandidate(
        filename="bad.pptx",
        local_path=tmp_path / "bad.pptx",
        data="BAD",
        local_hash="hash-bad",
    )
    ok_candidate = deck_pipeline_module._ExportCandidate(
        filename="ok.html",
        local_path=tmp_path / "ok.html",
        data="OK",
        local_hash="hash-ok",
    )
    monkeypatch.setattr(
        deck_pipeline_module,
        "_discover_export_files",
        lambda *args, **kwargs: [bad_candidate, ok_candidate],
    )
    transport = FakePushTransport(write_fails={"bad.pptx": TransportCallError("etag mismatch")})

    with pytest.raises(ExportConflictError, match="bad.pptx"):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    # Both writes were attempted -- the conflict did not stop the loop.
    written_paths = [c["files"][0]["path"] for c in transport.write_files_calls]
    assert written_paths == ["bad.pptx", "ok.html"]

    # Only the successful file's record landed; the conflicted file's own
    # state entry is untouched (absent, since it was never pushed before).
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags == {"export:ok.html": "hash-ok"}


def test_push_exports_conflict_preserves_an_already_recorded_unrelated_etag(tmp_path: Path, monkeypatch):
    """A conflict on one file must not disturb ANY other slug etag already
    on record -- including one belonging to a different artifact kind
    entirely (e.g. the prototype's own pull-side etag), proving `push_exports`
    never blindly overwrites `existing.etags` wholesale."""
    _seed_state(
        tmp_path,
        "pyforge-warden",
        etags={PROTOTYPE_ARTIFACT_KEY: "E-prototype-untouched"},
    )
    bad_candidate = deck_pipeline_module._ExportCandidate(
        filename="bad.pptx",
        local_path=tmp_path / "bad.pptx",
        data="BAD",
        local_hash="hash-bad",
    )
    monkeypatch.setattr(
        deck_pipeline_module,
        "_discover_export_files",
        lambda *args, **kwargs: [bad_candidate],
    )
    transport = FakePushTransport(write_fails={"bad.pptx": TransportCallError("etag mismatch")})

    with pytest.raises(ExportConflictError):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[PROTOTYPE_ARTIFACT_KEY] == "E-prototype-untouched"


def test_push_exports_auth_error_propagates_instead_of_being_treated_as_a_conflict(
    tmp_path: Path,
):
    """Regression: catching the base `errors.TransportError` conflated a
    genuine transport failure (expired credential, network outage) with a
    per-file conflict -- reporting it as "refused rather than risk
    clobbering a Design-side edit" and continuing to hammer the rest of the
    batch against a connection/credential that's still broken. `AuthError`
    (a `TransportError` subclass, not a `TransportCallError`) must
    propagate immediately instead."""
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    transport = FakePushTransport(write_fails={filename: AuthError("no credential")})

    with pytest.raises(AuthError):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)


# --- Story 23.4: PPTX discovery + binary write shape -------------------------


def test_discover_export_files_discovers_html_and_both_pptx_kinds(tmp_path: Path):
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    deck_bytes = b"DECK-BYTES"
    infographic_bytes = b"INFOGRAPHIC-BYTES"
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", deck_bytes)
    _write_export_pptx(tmp_path, "pyforge-warden", "infographic", "2026-08-07", infographic_bytes)

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    by_name = {c.filename: c for c in candidates}
    assert set(by_name) == {
        "pyforge-warden-infographic-standalone-2026-08-07.html",
        "pyforge-warden-deck-2026-08-07.pptx",
        "pyforge-warden_infographic_deck-2026-08-07.pptx",
    }
    assert by_name["pyforge-warden-infographic-standalone-2026-08-07.html"].binary is False
    assert by_name["pyforge-warden-deck-2026-08-07.pptx"].binary is True
    assert by_name["pyforge-warden_infographic_deck-2026-08-07.pptx"].binary is True


def test_discover_export_files_pptx_only_when_no_html_exists_yet(tmp_path: Path):
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", b"DECK")

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    assert len(candidates) == 1
    assert candidates[0].filename == "pyforge-warden-deck-2026-08-07.pptx"


def test_discover_export_files_picks_the_newest_pptx_of_each_kind(tmp_path: Path):
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-01", b"OLD")
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", b"NEW")

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    assert len(candidates) == 1
    assert base64.b64decode(candidates[0].data) == b"NEW"


def test_discover_export_files_pptx_candidate_is_base64_encoded_and_hashed_over_raw_bytes(
    tmp_path: Path,
):
    raw = b"RAW-PPTX-BYTES"
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", raw)

    candidates = deck_pipeline_module._discover_export_files(
        tmp_path / "presentations" / "pyforge-warden", "pyforge-warden"
    )

    assert len(candidates) == 1
    candidate = candidates[0]
    assert base64.b64decode(candidate.data) == raw
    assert candidate.local_hash == hashlib.sha256(raw).hexdigest()


def test_discover_export_files_raises_herald_error_when_html_cannot_be_read(
    tmp_path: Path,
):
    marp_dir = tmp_path / "presentations" / "pyforge-warden" / "src" / "marp"
    marp_dir.mkdir(parents=True)
    # A directory where a file is expected: `read_text` raises `OSError`
    # (`IsADirectoryError`), exercising the same catch a genuinely
    # unreadable file would hit.
    (marp_dir / "pyforge-warden-infographic-standalone-2026-08-07.html").mkdir()

    with pytest.raises(HeraldError, match="could not read"):
        deck_pipeline_module._discover_export_files(tmp_path / "presentations" / "pyforge-warden", "pyforge-warden")


def test_discover_export_files_raises_herald_error_when_pptx_cannot_be_read(
    tmp_path: Path,
):
    pptx_dir = tmp_path / "presentations" / "pyforge-warden" / "src" / "pptx"
    pptx_dir.mkdir(parents=True)
    (pptx_dir / "pyforge-warden-deck-2026-08-07.pptx").mkdir()

    with pytest.raises(HeraldError, match="could not read"):
        deck_pipeline_module._discover_export_files(tmp_path / "presentations" / "pyforge-warden", "pyforge-warden")


def test_push_exports_pptx_write_uses_base64_encoding(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    raw = b"\x50\x4b\x03\x04FAKE-PPTX-BYTES"
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", raw)
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    filename = "pyforge-warden-deck-2026-08-07.pptx"
    assert result.pushed == (filename,)
    write_call = transport.write_files_calls[0]
    assert write_call["files"][0]["path"] == filename
    assert write_call["files"][0]["encoding"] == "base64"
    assert base64.b64decode(write_call["files"][0]["data"]) == raw
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[f"export:{filename}"] == hashlib.sha256(raw).hexdigest()


def test_push_exports_html_write_still_carries_no_encoding_key(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport()

    push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    write_call = transport.write_files_calls[0]
    assert "encoding" not in write_call["files"][0]


def test_push_exports_pushes_all_three_artifact_kinds_together(tmp_path: Path):
    """AC1: a seeded deck with both PPTX files and the poster present and
    previously unpushed -- all three are written (the PPTX pair via the
    base64/``encoding`` shape) and ``state.py`` records an ``export:``
    content-hash entry for each."""
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    deck_bytes = b"DECK-BYTES"
    info_bytes = b"INFOGRAPHIC-BYTES"
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", deck_bytes)
    _write_export_pptx(tmp_path, "pyforge-warden", "infographic", "2026-08-07", info_bytes)
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    html_filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    deck_filename = "pyforge-warden-deck-2026-08-07.pptx"
    info_filename = "pyforge-warden_infographic_deck-2026-08-07.pptx"
    assert set(result.pushed) == {html_filename, deck_filename, info_filename}
    assert result.skipped == ()

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[f"export:{html_filename}"] == hashlib.sha256(b"<html>v1</html>").hexdigest()
    assert recorded.etags[f"export:{deck_filename}"] == hashlib.sha256(deck_bytes).hexdigest()
    assert recorded.etags[f"export:{info_filename}"] == hashlib.sha256(info_bytes).hexdigest()

    by_path = {c["files"][0]["path"]: c["files"][0] for c in transport.write_files_calls}
    assert "encoding" not in by_path[html_filename]
    assert by_path[deck_filename]["encoding"] == "base64"
    assert by_path[info_filename]["encoding"] == "base64"


def test_push_exports_second_push_with_no_changes_skips_all_three_kinds(
    tmp_path: Path,
):
    """AC2: the same deck pushed again immediately with no local changes --
    nothing is pushed and no ``write_files`` call is made, across all three
    artifact kinds at once."""
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    deck_bytes = b"DECK-BYTES"
    info_bytes = b"INFOGRAPHIC-BYTES"
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", deck_bytes)
    _write_export_pptx(tmp_path, "pyforge-warden", "infographic", "2026-08-07", info_bytes)
    html_filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    deck_filename = "pyforge-warden-deck-2026-08-07.pptx"
    info_filename = "pyforge-warden_infographic_deck-2026-08-07.pptx"
    _seed_state(
        tmp_path,
        "pyforge-warden",
        etags={
            f"export:{html_filename}": hashlib.sha256(b"<html>v1</html>").hexdigest(),
            f"export:{deck_filename}": hashlib.sha256(deck_bytes).hexdigest(),
            f"export:{info_filename}": hashlib.sha256(info_bytes).hexdigest(),
        },
    )
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result.pushed == ()
    assert set(result.skipped) == {html_filename, deck_filename, info_filename}
    assert transport.finalize_plan_calls == []
    assert transport.write_files_calls == []


# --- Story 23.4: --prove ------------------------------------------------------


def test_push_exports_prove_false_makes_no_fetch_call_and_leaves_readme_untouched(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path)

    assert result.proven == ()
    assert transport.fetch_rendered_bytes_calls == []
    assert readme_path.read_text(encoding="utf-8") == "# Deck\n"


def test_push_exports_prove_true_nothing_to_push_makes_no_fetch_call(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    assert result == ExportPushResult(slug="pyforge-warden", pushed=(), skipped=(), proven=())
    assert transport.fetch_rendered_bytes_calls == []


def test_push_exports_prove_true_skips_an_unchanged_file(tmp_path: Path):
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    stored_hash = hashlib.sha256(b"<html>v1</html>").hexdigest()
    _seed_state(tmp_path, "pyforge-warden", etags={f"export:{filename}": stored_hash})
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    transport = FakePushTransport()

    result = push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    assert result == ExportPushResult(slug="pyforge-warden", pushed=(), skipped=(filename,), proven=())
    assert transport.fetch_rendered_bytes_calls == []


def test_push_exports_prove_true_all_match_marks_proven_and_appends_ledger_row(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    transport = FakePushTransport(rendered_bytes={filename: b"<html>v1</html>"})

    result = push_exports(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prove=True,
        now=lambda: datetime(2026, 9, 16, tzinfo=timezone.utc),
    )

    assert result.proven == (filename,)
    assert transport.fetch_rendered_bytes_calls == [{"project_id": "p-1", "path": filename}]
    readme_text = readme_path.read_text(encoding="utf-8")
    assert "## Ledger — 2026-09-16 push-and-prove (spec-design-sync-loop CAP-6)" in readme_text
    assert filename in readme_text
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[f"export:{filename}"] == hashlib.sha256(b"<html>v1</html>").hexdigest()


def test_push_exports_prove_true_mismatch_raises_and_does_not_update_state_or_ledger(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    transport = FakePushTransport(rendered_bytes={filename: b"<html>DIFFERENT</html>"})

    with pytest.raises(ReadBackMismatchError, match=filename):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert f"export:{filename}" not in recorded.etags
    assert readme_path.read_text(encoding="utf-8") == "# Deck\n"


def test_push_exports_prove_true_mismatch_reverts_a_previously_recorded_etag(
    tmp_path: Path,
):
    """A mismatch on a file that *had* a prior record must restore that
    prior value, not merely drop the key -- the file was genuinely pushed
    with new bytes (the hash comparison already proved it changed), so
    leaving the old hash in place is what makes a retry see it as
    still-changed rather than falsely "up to date with the bad push"."""
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    _seed_state(tmp_path, "pyforge-warden", etags={f"export:{filename}": "stale-hash"})
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v2</html>")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    transport = FakePushTransport(
        plan=PlanHandle(plan_token="tok", base_etags={filename: "E9"}),
        rendered_bytes={filename: b"<html>WRONG</html>"},
    )

    with pytest.raises(ReadBackMismatchError):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags[f"export:{filename}"] == "stale-hash"


def test_push_exports_prove_true_one_mismatch_does_not_block_another_files_ledger_row(
    tmp_path: Path,
):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    _write_export_pptx(tmp_path, "pyforge-warden", "deck", "2026-08-07", b"GOOD")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    html_filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    pptx_filename = "pyforge-warden-deck-2026-08-07.pptx"
    transport = FakePushTransport(
        rendered_bytes={
            html_filename: b"<html>WRONG</html>",  # mismatch
            pptx_filename: b"GOOD",  # matches
        }
    )

    with pytest.raises(ReadBackMismatchError, match=html_filename):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert f"export:{html_filename}" not in recorded.etags
    assert f"export:{pptx_filename}" in recorded.etags
    readme_text = readme_path.read_text(encoding="utf-8")
    assert pptx_filename in readme_text
    assert html_filename not in readme_text
    # The Ledger row's byte count must be the candidate's true raw byte
    # count (`len(b"GOOD")` == 4), never the base64-encoded wire string's
    # length (8 chars) -- guards `ledger_rows.append` staying keyed off
    # `_candidate_raw_bytes`, not `candidate.data`.
    assert f"| {pptx_filename} | 4 | identical" in readme_text


def test_push_exports_prove_does_not_call_fetch_for_a_conflicted_file(tmp_path: Path):
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    transport = FakePushTransport(write_fails={filename: TransportCallError("etag mismatch")})

    with pytest.raises(ExportConflictError):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    assert transport.fetch_rendered_bytes_calls == []


def test_push_exports_prove_true_fetch_failure_is_treated_like_a_mismatch(
    tmp_path: Path,
):
    """A transient `fetch_rendered_bytes` failure (network/HTTP during the
    GET, not a byte mismatch) must not propagate raw or abort the batch: the
    write already landed, so it is folded into `mismatches` exactly like a
    genuine byte mismatch -- reverted `export:` entry, no Ledger row,
    surfaced via the batched `ReadBackMismatchError`."""
    _seed_state(tmp_path, "pyforge-warden")
    _write_export_html(tmp_path, "pyforge-warden", "2026-08-07", "<html>v1</html>")
    readme_path = tmp_path / "presentations" / "pyforge-warden" / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    filename = "pyforge-warden-infographic-standalone-2026-08-07.html"
    transport = FakePushTransport(fetch_fails={filename: TransportCallError("network blip")})

    with pytest.raises(ReadBackMismatchError, match=filename):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert f"export:{filename}" not in recorded.etags
    assert readme_path.read_text(encoding="utf-8") == "# Deck\n"


def test_push_exports_prove_true_conflict_on_one_file_does_not_suppress_another_files_ledger_row(
    tmp_path: Path, monkeypatch
):
    """A same-batch conflict on one file must never suppress a different
    file's already-earned Ledger row or state record: the conflicted file
    never reaches the prove loop at all (it is never added to `pushed`), so
    the clean file's proof and Ledger append proceed independently."""
    _seed_state(tmp_path, "pyforge-warden")
    bad_candidate = deck_pipeline_module._ExportCandidate(
        filename="bad.pptx",
        local_path=tmp_path / "bad.pptx",
        data=base64.b64encode(b"BAD").decode("ascii"),
        local_hash="hash-bad",
        binary=True,
    )
    ok_candidate = deck_pipeline_module._ExportCandidate(
        filename="ok.pptx",
        local_path=tmp_path / "ok.pptx",
        data=base64.b64encode(b"GOOD").decode("ascii"),
        local_hash="hash-ok",
        binary=True,
    )
    monkeypatch.setattr(
        deck_pipeline_module,
        "_discover_export_files",
        lambda *args, **kwargs: [bad_candidate, ok_candidate],
    )
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    deck_dir.mkdir(parents=True)
    readme_path = deck_dir / "README.md"
    readme_path.write_text("# Deck\n", encoding="utf-8")
    transport = FakePushTransport(
        write_fails={"bad.pptx": TransportCallError("etag mismatch")},
        rendered_bytes={"ok.pptx": b"GOOD"},
    )

    with pytest.raises(ExportConflictError, match="bad.pptx"):
        push_exports(transport, slug="pyforge-warden", repo_root=tmp_path, prove=True)

    readme_text = readme_path.read_text(encoding="utf-8")
    assert "ok.pptx" in readme_text
    assert "bad.pptx" not in readme_text
    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "pyforge-warden")
    assert recorded.etags == {"export:ok.pptx": "hash-ok"}


# --- Story 23.4: _strip_serve_harness -----------------------------------------


def test_strip_serve_harness_removes_a_single_injected_style_tag():
    content = (
        b"<html><head><style data-omelette-injected>body{color:red}</style>"
        b"<title>x</title></head><body>hi</body></html>"
    )

    stripped = deck_pipeline_module._strip_serve_harness(content, path="poster.html")

    assert b"data-omelette-injected" not in stripped
    assert stripped == b"<html><head>\n<title>x</title></head><body>hi</body></html>"


def test_strip_serve_harness_removes_multiple_contiguous_tags():
    content = (
        b"<head><style data-omelette-injected>a</style><script data-omelette-injected>b</script><title>t</title></head>"
    )

    stripped = deck_pipeline_module._strip_serve_harness(content, path="x.html")

    assert stripped == b"<head>\n<title>t</title></head>"


def test_strip_serve_harness_tolerates_whitespace_between_tags():
    content = (
        b"<head>\n  <style data-omelette-injected>a</style>\n\n"
        b"<script data-omelette-injected>b</script>\n<title>t</title></head>"
    )

    stripped = deck_pipeline_module._strip_serve_harness(content, path="x.html")

    assert b"data-omelette-injected" not in stripped
    assert stripped.startswith(b"<head>\n")
    assert stripped.endswith(b"<title>t</title></head>")


def test_strip_serve_harness_passthrough_for_a_non_html_path():
    content = b"\x50\x4b\x03\x04-fake-pptx-bytes-data-omelette-injected"

    assert deck_pipeline_module._strip_serve_harness(content, path="a.pptx") == content


def test_strip_serve_harness_no_head_tag_returns_content_unchanged():
    content = b"<html><body>no head here</body></html>"

    assert deck_pipeline_module._strip_serve_harness(content, path="x.html") == content


def test_strip_serve_harness_head_with_nothing_injected_returns_content_unchanged():
    content = b"<head><title>t</title></head>"

    assert deck_pipeline_module._strip_serve_harness(content, path="x.html") == content


def test_strip_serve_harness_ignores_a_style_tag_without_the_marker():
    content = b"<head><style>ordinary</style><title>t</title></head>"

    assert deck_pipeline_module._strip_serve_harness(content, path="x.html") == content


# --- Story 23.3: PptxTemplateExporter / select_exporter ----------------------


def _init_git_repo(root: Path) -> None:
    subprocess.run(["git", "init", "-q"], cwd=root, check=True)
    subprocess.run(["git", "config", "user.email", "test@example.com"], cwd=root, check=True)
    subprocess.run(["git", "config", "user.name", "Test"], cwd=root, check=True)
    (root / "README.md").write_text("scratch repo\n", encoding="utf-8")
    subprocess.run(["git", "add", "-A"], cwd=root, check=True)
    subprocess.run(["git", "commit", "-q", "-m", "init"], cwd=root, check=True)


def _make_potx_deck(tmp_path: Path, slug: str, *, template_rel: str) -> Path:
    """A deck directory with a real (blank) ``.pptx`` template at
    ``template_rel``, a README declaring it via ``register_potx_template``,
    and a minimal valid ``content_plan.json`` -- the shape
    ``PptxTemplateExporter``/``select_exporter`` need to run for real,
    without mocking ``pptx_pipeline`` itself (it is already covered by
    ``test_pptx_pipeline.py``)."""
    deck_dir = tmp_path / "presentations" / slug
    deck_dir.mkdir(parents=True)
    (deck_dir / "README.md").write_text(f"# {slug}\n\nBody.\n", encoding="utf-8")
    template_path = tmp_path / template_rel
    template_path.parent.mkdir(parents=True, exist_ok=True)
    Presentation().save(str(template_path))
    registry.register_potx_template(deck_dir / "README.md", template_rel)
    (deck_dir / "src").mkdir(parents=True, exist_ok=True)
    (deck_dir / "src" / "content_plan.json").write_text(
        json.dumps({"slides": [{"layout": 0, "placeholders": {}}]}), encoding="utf-8"
    )
    return deck_dir


def test_select_exporter_returns_pixi_exporter_when_no_potx_declared(tmp_path: Path):
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    deck_dir.mkdir(parents=True)
    (deck_dir / "README.md").write_text("# pyforge-warden\n", encoding="utf-8")

    exporter = select_exporter(slug="pyforge-warden", repo_root=tmp_path)

    assert isinstance(exporter, PixiDeckExporter)


def test_select_exporter_returns_pptx_template_exporter_when_potx_declared(
    tmp_path: Path,
):
    _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )

    exporter = select_exporter(slug="pyforge-warden", repo_root=tmp_path)

    assert isinstance(exporter, PptxTemplateExporter)


def test_pull_prototype_with_no_explicit_exporter_routes_through_select_exporter(tmp_path: Path, monkeypatch):
    """Regression proof for the 3 call-site swap: with no injected
    ``exporter=``, ``pull_prototype`` must resolve via ``select_exporter``
    (and therefore honor a declared ``.potx`` template) rather than always
    defaulting straight to ``PixiDeckExporter``."""
    deck_dir = _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )
    _init_git_repo(tmp_path)
    _seed_state(tmp_path, "pyforge-warden", etags={PROTOTYPE_ARTIFACT_KEY: "old"})
    transport = FakePullTransport(
        answers=FileRead(path="PyForge Warden.dc.html", etag="new", body="<html>v2</html>", unchanged=False)
    )
    calls: list[str] = []
    monkeypatch.setattr(
        PptxTemplateExporter,
        "export",
        lambda self, *, slug, repo_root: calls.append(slug),
    )

    pull_prototype(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        prover=FakeProver(),
    )

    assert calls == ["pyforge-warden"]
    assert (deck_dir / "project" / "PyForge Warden.dc.html").read_text(encoding="utf-8") == "<html>v2</html>"


def test_pull_marp_source_with_no_explicit_exporter_routes_through_select_exporter(tmp_path: Path, monkeypatch):
    """Regression proof for the 3 call-site swap (mirrors
    ``test_pull_prototype_with_no_explicit_exporter_routes_through_select_exporter``):
    with no injected ``exporter=``, ``pull_marp_source`` must resolve via
    ``select_exporter`` (and therefore honor a declared ``.potx`` template)
    rather than always defaulting straight to ``PixiDeckExporter``. Without
    this test, a future revert of just this call site back to a hardcoded
    ``PixiDeckExporter()`` would pass every existing test."""
    _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )
    _init_git_repo(tmp_path)
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="new", body="# Deck", unchanged=False))
    calls: list[str] = []
    monkeypatch.setattr(
        PptxTemplateExporter,
        "export",
        lambda self, *, slug, repo_root: calls.append(slug),
    )

    pull_marp_source(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        kind="deck",
        now=lambda: _FIXED_NOW,
    )

    assert calls == ["pyforge-warden"]


def test_pull_standalone_bundle_with_no_explicit_exporter_routes_through_select_exporter(tmp_path: Path, monkeypatch):
    """Regression proof for the 3 call-site swap (mirrors
    ``test_pull_prototype_with_no_explicit_exporter_routes_through_select_exporter``):
    with no injected ``exporter=``, ``pull_standalone_bundle`` must resolve
    via ``select_exporter`` (and therefore honor a declared ``.potx``
    template) rather than always defaulting straight to
    ``PixiDeckExporter``. Without this test, a future revert of just this
    call site back to a hardcoded ``PixiDeckExporter()`` would pass every
    existing test."""
    _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )
    _init_git_repo(tmp_path)
    _seed_state(tmp_path, "pyforge-warden")
    transport = FakePullTransport(answers=FileRead(path="x", etag="new", body="<html>bundle</html>", unchanged=False))
    calls: list[str] = []
    monkeypatch.setattr(
        PptxTemplateExporter,
        "export",
        lambda self, *, slug, repo_root: calls.append(slug),
    )

    pull_standalone_bundle(
        transport,
        slug="pyforge-warden",
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )

    assert calls == ["pyforge-warden"]


def test_pptx_template_exporter_fills_the_template_and_stamps_the_pptx(
    tmp_path: Path,
):
    deck_dir = _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )
    _init_git_repo(tmp_path)
    html_exporter = FakeExporter()

    PptxTemplateExporter(html_exporter=html_exporter, now=lambda: datetime(2026, 9, 16, tzinfo=timezone.utc)).export(
        slug="pyforge-warden", repo_root=tmp_path
    )

    out_path = deck_dir / "src" / "pptx" / "pyforge-warden-deck-2026-09-16.pptx"
    assert out_path.is_file()
    # A real, filled presentation -- not a Marp render.
    prs = Presentation(str(out_path))
    assert len(prs.slides) == 1
    # html/infographic-pptx still produced via deck-export, explicit
    # targets excluding deck-pptx.
    assert html_exporter.calls == [("pyforge-warden", tmp_path)]
    # The filled PPTX is stamped; deck-export's own subprocess is
    # responsible for stamping html/infographic-pptx (out of this class's
    # scope).
    assert stamps.read_stamp(out_path) is not None


def test_pptx_template_exporter_raises_when_no_potx_registered(tmp_path: Path):
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    deck_dir.mkdir(parents=True)
    (deck_dir / "README.md").write_text("# pyforge-warden\n", encoding="utf-8")

    with pytest.raises(HeraldError, match="no PowerPoint template registered"):
        PptxTemplateExporter(html_exporter=FakeExporter()).export(slug="pyforge-warden", repo_root=tmp_path)


def test_pptx_template_exporter_raises_naming_the_missing_content_plan(tmp_path: Path):
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    deck_dir.mkdir(parents=True)
    (deck_dir / "README.md").write_text("# pyforge-warden\n", encoding="utf-8")
    template_path = deck_dir / "project" / "deck.pptx"
    template_path.parent.mkdir(parents=True)
    Presentation().save(str(template_path))
    registry.register_potx_template(deck_dir / "README.md", "presentations/pyforge-warden/project/deck.pptx")

    content_plan_path = deck_dir / "src" / "content_plan.json"
    with pytest.raises(HeraldError, match=str(content_plan_path)):
        PptxTemplateExporter(html_exporter=FakeExporter()).export(slug="pyforge-warden", repo_root=tmp_path)


def test_pptx_template_exporter_propagates_a_missing_template_as_pptx_template_error(
    tmp_path: Path,
):
    """The template-file-missing case propagates unchanged from
    ``pptx_pipeline.run_fill`` (unlike the content-plan case above, which
    this class pre-checks itself)."""
    deck_dir = tmp_path / "presentations" / "pyforge-warden"
    deck_dir.mkdir(parents=True)
    (deck_dir / "README.md").write_text("# pyforge-warden\n", encoding="utf-8")
    registry.register_potx_template(deck_dir / "README.md", "presentations/pyforge-warden/project/missing.pptx")
    (deck_dir / "src").mkdir(parents=True)
    (deck_dir / "src" / "content_plan.json").write_text(json.dumps({"slides": []}), encoding="utf-8")

    with pytest.raises(PptxTemplateError):
        PptxTemplateExporter(html_exporter=FakeExporter()).export(slug="pyforge-warden", repo_root=tmp_path)


def test_pptx_template_exporter_propagates_an_html_exporter_failure(tmp_path: Path):
    _make_potx_deck(
        tmp_path,
        "pyforge-warden",
        template_rel="presentations/pyforge-warden/project/deck.pptx",
    )
    html_exporter = FakeExporter(fails=HeraldError("deck-export failed"))

    with pytest.raises(HeraldError, match="deck-export failed"):
        PptxTemplateExporter(html_exporter=html_exporter).export(slug="pyforge-warden", repo_root=tmp_path)


def test_pixi_partial_deck_exporter_excludes_deck_pptx_from_its_subprocess_command(monkeypatch, tmp_path: Path):
    """The one hardcoded invariant this class exists for: its shelled
    ``deck-export`` targets must be exactly ``html``/``infographic-pptx``,
    never ``deck-pptx`` -- that target is what ``PptxTemplateExporter``
    fills itself. A regression here would silently overwrite a just-filled
    ``.potx``-based PowerPoint with a Marp render (Story 23.3 follow-up
    review)."""
    calls: list[list[str]] = []

    class _Completed:
        returncode = 0
        stdout = ""
        stderr = ""

    def _run(cmd, **kwargs):
        calls.append(cmd)
        return _Completed()

    monkeypatch.setattr(subprocess, "run", _run)
    _PixiPartialDeckExporter().export(slug="pyforge-warden", repo_root=tmp_path)

    assert calls == [
        ["pixi", "run", "-e", "pyforge-guild", "deck-export", "pyforge-warden", "html", "infographic-pptx"]
    ]
    assert "deck-pptx" not in calls[0]


# --- Story 23.2 (CAP-2 from spec-design-sync-loop): _windowed_read ----------


def test_windowed_read_short_circuits_on_unchanged():
    transport = FakePullTransport(answers=FileRead(path="x", etag="E1", body=None, unchanged=True))
    result = _windowed_read(transport, project_id="p-1", path="x.dc.html", if_none_match="E1")
    assert result.unchanged is True
    assert transport.calls == [{"project_id": "p-1", "path": "x.dc.html", "if_none_match": "E1"}]


def test_windowed_read_single_call_when_the_server_answers_whole():
    transport = FakePullTransport(answers=FileRead(path="x", etag="E2", body="<html></html>", unchanged=False))
    result = _windowed_read(transport, project_id="p-1", path="x.dc.html")
    assert result == FileRead(path="x.dc.html", etag="E2", body="<html></html>", unchanged=False)
    assert len(transport.calls) == 1


def test_windowed_read_self_short_circuits_when_the_server_ignores_if_none_match():
    """Live-verified 2026-09-18: the server never answers its own
    ``{unchanged: true}`` short-circuit once a file needs windowing --
    an ``if_none_match`` that exactly matches the current etag still
    comes back as an ordinary truncated first window. Without this
    function's own etag check, a large unchanged file would be
    re-downloaded (and re-written) on every single run, breaking CAP-30's
    idempotency requirement for exactly the artifact that motivated
    ``_windowed_read`` in the first place -- caught live pulling
    ``six-quarter-roadmap``'s 320 KB prototype a second time."""
    transport = FakePullTransport(
        answers=[
            FileRead(
                path="x",
                etag="E1",
                body="line1\nline2773",
                unchanged=False,  # the server itself never sets this here
                first_line=1,
                last_line=2773,
                total_lines=3377,
            )
        ]
    )
    result = _windowed_read(transport, project_id="p-1", path="big.dc.html", if_none_match="E1")
    assert result.unchanged is True
    assert result.etag == "E1"
    assert len(transport.calls) == 1  # never asked for a second window


def test_windowed_read_reassembles_across_multiple_windows():
    """Mirrors the live-verified 2026-09-18 evidence: a 3377-line file
    answered as two windows (``1-2773``/``total_lines=3377``, then
    ``2774-3377``), joined back into one file with exactly the newline
    that separated line 2773 from line 2774 in the original -- neither
    doubled nor lost."""
    transport = FakePullTransport(
        answers=[
            FileRead(
                path="x",
                etag="E3",
                body="line1\nline2773",
                unchanged=False,
                first_line=1,
                last_line=2773,
                total_lines=3377,
            ),
            FileRead(
                path="x",
                etag="E3",
                body="line2774\nline3377",
                unchanged=False,
                first_line=2774,
                last_line=3377,
                total_lines=3377,
            ),
        ]
    )
    result = _windowed_read(transport, project_id="p-1", path="big.dc.html")
    assert result.unchanged is False
    assert result.etag == "E3"
    assert result.body == "line1\nline2773\nline2774\nline3377"
    assert len(transport.calls) == 2
    assert transport.calls[1] == {
        "project_id": "p-1",
        "path": "big.dc.html",
        "offset": 2774,
    }


def test_windowed_read_strips_the_truncation_trailer_from_a_non_final_window():
    """Live-verified 2026-09-18 (see ``_TRUNCATION_TRAILER_RE``'s own
    docstring): a truncated window's raw ``body`` carries two bogus extra
    lines -- a blank separator, then a resumption hint -- beyond its own
    declared ``lines="A-B"`` span. Left unstripped, the reconstructed file
    gains two lines that were never in the original (caught live pulling
    ``six-quarter-roadmap``'s 320 KB prototype: 3379 lines landed locally
    against a real ``total_lines=3377``)."""
    transport = FakePullTransport(
        answers=[
            FileRead(
                path="x",
                etag="E6",
                body=(
                    "line1\nline2773\n\n…[+58491 bytes truncated at "
                    "read_file's 256 KiB cap — the body ends at a complete "
                    "line; continue with offset=2774]"
                ),
                unchanged=False,
                first_line=1,
                last_line=2773,
                total_lines=3377,
            ),
            FileRead(
                path="x",
                etag="E6",
                body="line2774\nline3377",
                unchanged=False,
                first_line=2774,
                last_line=3377,
                total_lines=3377,
            ),
        ]
    )
    result = _windowed_read(transport, project_id="p-1", path="big.dc.html")
    assert result.body == "line1\nline2773\nline2774\nline3377"
    assert result.body.count("\n") == 3  # 4 real lines, no bogus ones


def test_windowed_read_refuses_a_changed_answer_with_no_body():
    transport = FakePullTransport(answers=FileRead(path="x", etag="E4", body=None, unchanged=False))
    with pytest.raises(HeraldError, match="returned no body"):
        _windowed_read(transport, project_id="p-1", path="x.dc.html")


def test_windowed_read_refuses_a_partial_window_with_no_resume_point():
    transport = FakePullTransport(
        answers=FileRead(
            path="x",
            etag="E5",
            body="partial",
            unchanged=False,
            first_line=1,
            last_line=None,
            total_lines=None,
        )
    )
    # first_line is set with no last_line/total_lines -- a malformed,
    # unresumable partial window (see FileRead.truncated's own docstring
    # for why this combination is a genuine "no window at all" only when
    # BOTH are None).
    with pytest.raises(HeraldError, match="no last_line/total_lines"):
        _windowed_read(transport, project_id="p-1", path="x.dc.html")


def test_windowed_read_raises_pagination_stalled_error_on_a_non_advancing_window():
    """Story 24.1 (DW-FU-23-2): the un-guarded loop broke only on
    ``last_line >= total_lines`` and never checked that ``last_line``
    actually advanced between calls, so a server that kept answering the
    window it was paged FOR with the same ``(last_line, total_lines)``
    pair would loop forever. A fake transport reproducing exactly that --
    the second window ends at the same line the first one did -- must
    raise ``PaginationStalledError`` naming the file and the stalled line,
    never loop or silently truncate."""
    transport = FakePullTransport(
        answers=[
            FileRead(
                path="x",
                etag="E7",
                body="line1\nline2773",
                unchanged=False,
                first_line=1,
                last_line=2773,
                total_lines=6000,
            ),
            FileRead(
                path="x",
                etag="E7",
                body="line1\nline2773",
                unchanged=False,
                first_line=1,
                last_line=2773,  # stalled -- did not advance past window 1
                total_lines=6000,
            ),
        ]
    )
    with pytest.raises(PaginationStalledError, match=r"big\.dc\.html.*line 2773.*line 2773"):
        _windowed_read(transport, project_id="p-1", path="big.dc.html")
    # Refused after exactly the second (stalled) window -- never asked for
    # a third, and never fell back to looping.
    assert len(transport.calls) == 2


def test_windowed_read_raises_pagination_stalled_error_on_a_regressed_window():
    """The guard's ``<=`` comparison also refuses a window whose
    ``last_line`` regresses below the previous window's, not just an exact
    repeat -- a second, distinct failure shape covered separately from the
    exact-equality case above."""
    transport = FakePullTransport(
        answers=[
            FileRead(
                path="x",
                etag="E8",
                body="line1\nline2773",
                unchanged=False,
                first_line=1,
                last_line=2773,
                total_lines=6000,
            ),
            FileRead(
                path="x",
                etag="E8",
                body="line1\nline1000",
                unchanged=False,
                first_line=1,
                last_line=1000,  # regressed -- below window 1's last_line
                total_lines=6000,
            ),
        ]
    )
    with pytest.raises(PaginationStalledError, match=r"big\.dc\.html.*line 1000.*line 2773"):
        _windowed_read(transport, project_id="p-1", path="big.dc.html")
    assert len(transport.calls) == 2


# --- Story 23.2 (CAP-2 from spec-design-sync-loop): adopt --------------------


def test_adopt_bootstraps_a_new_twin_and_registers(tmp_path: Path):
    transport = FakePullTransport(
        answers=[
            FileRead(path="x", etag="E1", body="<html>one</html>", unchanged=False),
            FileRead(path="y", etag="E2", body="<html>two</html>", unchanged=False),
        ]
    )
    dest_dir = tmp_path / "presentations" / "six-quarter-roadmap"

    result = adopt(
        transport,
        state_key="six-quarter-roadmap",
        project_id="p-roadmap",
        project_name="PyForge six-quarter roadmap",
        project_url="https://claude.ai/design/p/p-roadmap",
        artifacts=[
            ("PyForge Roadmap.dc.html", "project/PyForge Roadmap.dc.html"),
            ("github.md", "project/github.md"),
        ],
        dest_dir=dest_dir,
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )

    assert result.bootstrapped is True
    assert result.registered is True
    assert result.artifacts == (
        AdoptedArtifact(
            remote_path="PyForge Roadmap.dc.html",
            local_path=dest_dir / "project" / "PyForge Roadmap.dc.html",
            unchanged=False,
        ),
        AdoptedArtifact(
            remote_path="github.md",
            local_path=dest_dir / "project" / "github.md",
            unchanged=False,
        ),
    )
    assert (dest_dir / "project" / "PyForge Roadmap.dc.html").read_text(encoding="utf-8") == "<html>one</html>"
    assert (dest_dir / "project" / "github.md").read_text(encoding="utf-8") == "<html>two</html>"

    registered = read_registry(dest_dir / "README.md")
    assert registered.project_name == "PyForge six-quarter roadmap"
    assert registered.project_id == "p-roadmap"
    assert registered.file_url == "https://claude.ai/design/p/p-roadmap"

    recorded = state.read(tmp_path / state.DEFAULT_STATE_PATH, "six-quarter-roadmap")
    assert recorded.project_id == "p-roadmap"
    assert recorded.etags == {"PyForge Roadmap.dc.html": "E1", "github.md": "E2"}
    assert recorded.last_pull == _FIXED_NOW.isoformat()


def test_adopt_second_call_against_unchanged_content_writes_nothing(tmp_path: Path):
    dest_dir = tmp_path / "presentations" / "six-quarter-roadmap"
    first_transport = FakePullTransport(
        answers=[FileRead(path="x", etag="E1", body="<html>one</html>", unchanged=False)]
    )
    adopt(
        first_transport,
        state_key="six-quarter-roadmap",
        project_id="p-roadmap",
        project_name="PyForge six-quarter roadmap",
        project_url="https://claude.ai/design/p/p-roadmap",
        artifacts=[("PyForge Roadmap.dc.html", "project/PyForge Roadmap.dc.html")],
        dest_dir=dest_dir,
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )
    readme_before = (dest_dir / "README.md").read_text(encoding="utf-8")

    second_transport = FakePullTransport(answers=[FileRead(path="x", etag="E1", body=None, unchanged=True)])
    result = adopt(
        second_transport,
        state_key="six-quarter-roadmap",
        project_id="p-roadmap",
        project_name="PyForge six-quarter roadmap",
        project_url="https://claude.ai/design/p/p-roadmap",
        artifacts=[("PyForge Roadmap.dc.html", "project/PyForge Roadmap.dc.html")],
        dest_dir=dest_dir,
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )

    assert result.bootstrapped is False
    assert result.registered is False
    assert result.artifacts == (
        AdoptedArtifact(
            remote_path="PyForge Roadmap.dc.html",
            local_path=dest_dir / "project" / "PyForge Roadmap.dc.html",
            unchanged=True,
        ),
    )
    assert second_transport.calls[0]["if_none_match"] == "E1"
    # The registry section is byte-identical -- adopt never rewrites it once
    # a section already parses.
    assert (dest_dir / "README.md").read_text(encoding="utf-8") == readme_before


def test_adopt_skips_the_registry_for_a_design_system_mirror(tmp_path: Path):
    transport = FakePullTransport(answers=[FileRead(path="x", etag="E1", body="body { color: red; }", unchanged=False)])
    dest_dir = tmp_path / "presentations" / "_design-systems" / "modernist"

    result = adopt(
        transport,
        state_key="design-system-modernist",
        project_id="p-modernist",
        artifacts=[("styles.css", "styles.css")],
        dest_dir=dest_dir,
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )

    assert result.registered is False
    assert not (dest_dir / "README.md").exists()
    assert (dest_dir / "styles.css").read_text(encoding="utf-8") == "body { color: red; }"


def test_adopt_refuses_an_empty_artifact_list(tmp_path: Path):
    with pytest.raises(HeraldError, match="no artifacts given"):
        adopt(
            FakePullTransport(answers=None),
            state_key="six-quarter-roadmap",
            project_id="p-roadmap",
            artifacts=[],
            dest_dir=tmp_path / "presentations" / "six-quarter-roadmap",
            repo_root=tmp_path,
        )


def test_adopt_refuses_a_project_id_mismatch_on_a_later_call(tmp_path: Path):
    dest_dir = tmp_path / "presentations" / "six-quarter-roadmap"
    adopt(
        FakePullTransport(answers=[FileRead(path="x", etag="E1", body="one", unchanged=False)]),
        state_key="six-quarter-roadmap",
        project_id="p-roadmap",
        artifacts=[("a.md", "a.md")],
        dest_dir=dest_dir,
        repo_root=tmp_path,
        now=lambda: _FIXED_NOW,
    )
    with pytest.raises(HeraldError, match="already adopted against Design project"):
        adopt(
            FakePullTransport(answers=None),
            state_key="six-quarter-roadmap",
            project_id="p-different",
            artifacts=[("a.md", "a.md")],
            dest_dir=dest_dir,
            repo_root=tmp_path,
        )
