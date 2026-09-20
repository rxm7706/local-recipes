"""Story 15.1, CAP-1: template-parse-then-fill produces a genuinely
editable deck.

Covers the story spec's I/O & Edge-Case Matrix directly:
``extract_spec``/``fill_template`` validation, the ``run_spec``/``run_fill``
file-path wrappers, the CLI's ``pptx-spec``/``pptx-fill`` subcommands, and
the flagship round-trip proof (real ``<a:t>`` runs, no ``<p:pic>``, and an
independent ``soffice --headless --convert-to pptx`` round-trip) using 3
real slides transcribed from
``presentations/pyforge-herald/src/marp/pyforge-herald-deck-2026-07-24.md``.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from pyforge.herald import cli, errors, pptx_pipeline

# === shared fixtures/helpers =================================================


@pytest.fixture(scope="module")
def template_path() -> Path:
    return pptx_pipeline.default_template_path()


def _write_content_plan(tmp_path: Path, plan: dict) -> Path:
    path = tmp_path / "content_plan.json"
    path.write_text(json.dumps(plan), encoding="utf-8")
    return path


def _slide_xml(pptx_path: Path, slide_number: int) -> str:
    with zipfile.ZipFile(pptx_path) as archive:
        return archive.read(f"ppt/slides/slide{slide_number}.xml").decode("utf-8")


class _FakePlaceholderFormat:
    """Stand-in for python-pptx's ``PlaceholderFormat`` -- just the two
    attributes ``extract_spec``/``_resolve_placeholder_values`` read."""

    def __init__(self, idx: int, type_: object) -> None:
        self.idx = idx
        self.type = type_


class _FakePlaceholder:
    """Stand-in for a python-pptx layout placeholder shape, minimal enough
    to drive ``extract_spec``'s geometry/type fallbacks and
    ``_resolve_placeholder_values``'s ``has_text_frame`` guard without a
    real ``.pptx`` -- neither branch is reachable via the bundled default
    template (every one of its placeholders has an explicit type, explicit
    geometry, and ``has_text_frame=True``)."""

    def __init__(
        self,
        idx: int,
        type_: object = None,
        name: str = "Fake Placeholder",
        has_text_frame: bool = True,
        left: int | None = None,
        top: int | None = None,
        width: int | None = None,
        height: int | None = None,
    ) -> None:
        self.placeholder_format = _FakePlaceholderFormat(idx, type_)
        self.name = name
        self.has_text_frame = has_text_frame
        self.left = left
        self.top = top
        self.width = width
        self.height = height


class _FakeLayout:
    """Stand-in for a python-pptx ``SlideLayout``: just ``name`` +
    ``placeholders``, the two attributes ``extract_spec``/
    ``_resolve_layout``/``_resolve_placeholder_values`` read."""

    def __init__(self, name: str, placeholders: list[_FakePlaceholder]) -> None:
        self.name = name
        self.placeholders = placeholders


class _FakePresentation:
    """Stand-in for python-pptx's ``Presentation``: just ``slide_layouts``,
    swapped in for ``_open_template``'s real return value via monkeypatch
    so ``extract_spec``/``fill_template`` run their real logic against a
    fake in-memory template."""

    def __init__(self, layouts: list[_FakeLayout]) -> None:
        self.slide_layouts = layouts


# The 3 real slides (cover, "Invisible engineering is failed engineering",
# "The Design-Code bridge -- realized" bullets), transcribed verbatim (minus
# markdown ** bold markers) from
# presentations/pyforge-herald/src/marp/pyforge-herald-deck-2026-07-24.md.
_THREE_REAL_SLIDES_PLAN = {
    "slides": [
        {
            "layout": 0,  # Title Slide
            "placeholders": {
                "0": "capture the dream. proclaim the release.",
                "1": (
                    "Herald is the visual media and communications engine "
                    "of the Dream-to-Code factory — the first to touch the "
                    "Dream and the last voice in the pipeline. Its first "
                    "organ, the Design↔Code bridge, is already real."
                ),
            },
        },
        {
            "layout": 1,  # Title and Content
            "placeholders": {
                "0": "Invisible engineering is failed engineering",
                "1": (
                    "Herald rejects dry, unreadable raw logs — pipeline "
                    "telemetry becomes scannable, visually striking "
                    "dashboards and briefings that reflect the original "
                    "vision."
                ),
            },
        },
        {
            "layout": 1,  # Title and Content
            "placeholders": {
                "0": "The Design–Code bridge — realized",
                "1": [
                    "seed a Design project from the repo — Modernist-bound, runtime included",
                    "design visually in Claude Design",
                    "pull the prototype back — extract → build → export, zero downloads",
                    "etags catch conflicts; the loop is specced as the herald CLI (5 CAPs)",
                ],
            },
        },
    ]
}


# === extract_spec / spec_to_dict =============================================


def test_emu_normalizes_none_to_zero():
    assert pptx_pipeline._emu(None) == 0


def test_emu_passes_through_an_explicit_value():
    assert pptx_pipeline._emu(12345) == 12345


def test_extract_spec_unset_placeholder_type_falls_back_to_unknown(monkeypatch):
    """python-pptx returns ``None`` for ``placeholder_format.type`` only
    when a ``<p:ph>`` element carries no ``type`` attribute at all --
    unreachable via the bundled default template (every one of its
    placeholders has an explicit type), so this pins the ``"UNKNOWN"``
    fallback with a constructed fake layout instead."""
    fake_prs = _FakePresentation([_FakeLayout("Fake Layout", [_FakePlaceholder(idx=0, type_=None)])])
    monkeypatch.setattr(pptx_pipeline, "_open_template", lambda path: fake_prs)

    spec = pptx_pipeline.extract_spec(Path("unused-template.pptx"))

    assert spec.layouts[0].placeholders[0].type == "UNKNOWN"


def test_extract_spec_matches_python_pptx_for_every_layout(template_path: Path):
    spec = pptx_pipeline.extract_spec(template_path)
    prs = Presentation(str(template_path))

    assert len(spec.layouts) == len(prs.slide_layouts)
    for layout_spec, layout in zip(spec.layouts, prs.slide_layouts, strict=True):
        assert layout_spec.name == layout.name
        # `LayoutPlaceholders.__getitem__` is positional (list index), NOT
        # keyed by the placeholder's own `idx` attribute -- unlike
        # `SlidePlaceholders.__getitem__`, which fill_template relies on.
        # Build our own idx-keyed map here rather than indexing by idx.
        by_idx = {ph.placeholder_format.idx: ph for ph in layout.placeholders}
        actual_idxs = {p.idx for p in layout_spec.placeholders}
        assert actual_idxs == set(by_idx)
        for ph_spec in layout_spec.placeholders:
            ph = by_idx[ph_spec.idx]
            assert ph_spec.type == ph.placeholder_format.type.name
            assert ph_spec.name == ph.name
            assert ph_spec.left == ph.left
            assert ph_spec.top == ph.top
            assert ph_spec.width == ph.width
            assert ph_spec.height == ph.height


def test_extract_spec_raises_pptx_template_error_when_template_missing(tmp_path: Path):
    with pytest.raises(errors.PptxTemplateError):
        pptx_pipeline.extract_spec(tmp_path / "does-not-exist.pptx")


def test_extract_spec_raises_pptx_template_error_on_unopenable_file(tmp_path: Path):
    bad = tmp_path / "not-a-pptx.pptx"
    bad.write_text("not a zip file at all", encoding="utf-8")
    with pytest.raises(errors.PptxTemplateError):
        pptx_pipeline.extract_spec(bad)


def test_spec_to_dict_round_trips_as_json(template_path: Path):
    spec = pptx_pipeline.extract_spec(template_path)
    parsed = json.loads(json.dumps(pptx_pipeline.spec_to_dict(spec)))
    assert len(parsed["layouts"]) == len(spec.layouts)


# === fill_template validation =================================================


def test_fill_template_unknown_layout_index_raises(template_path: Path):
    plan = {"slides": [{"layout": 999, "placeholders": {}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_unknown_layout_name_raises(template_path: Path):
    plan = {"slides": [{"layout": "Not A Real Layout", "placeholders": {}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_unknown_placeholder_idx_raises(template_path: Path):
    plan = {"slides": [{"layout": 0, "placeholders": {"99": "no such placeholder"}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


@pytest.mark.parametrize("value", [123, None, {"nested": "dict"}, [], [1, 2], ["ok", 3]])
def test_fill_template_malformed_placeholder_value_raises(template_path: Path, value):
    plan = {"slides": [{"layout": 0, "placeholders": {"0": value}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_non_text_placeholder_raises_invalid_content_plan_error(
    monkeypatch,
):
    """A placeholder whose ``has_text_frame`` is ``False`` must raise
    ``InvalidContentPlanError`` cleanly during validation, not crash
    ``_set_placeholder_text`` with an unhandled exception -- validation
    happens before any slide is added (I/O matrix's "No file written"
    invariant). Unreachable via the bundled default template as generated
    by python-pptx 1.0.2: even its "Picture with Caption" layout's idx-1
    ``PICTURE`` placeholder reports ``has_text_frame=True`` while unfilled
    (an empty picture placeholder is backed by a plain ``<p:sp>`` element,
    identical in kind to a text placeholder, until an image is actually
    inserted via ``insert_picture()``) -- so this pins the guard with a
    constructed fake layout whose placeholder genuinely cannot hold text."""
    fake_prs = _FakePresentation(
        [
            _FakeLayout(
                "Fake Picture Layout",
                [_FakePlaceholder(idx=1, type_="PICTURE", has_text_frame=False)],
            )
        ]
    )
    monkeypatch.setattr(pptx_pipeline, "_open_template", lambda path: fake_prs)

    plan = {"slides": [{"layout": 0, "placeholders": {"1": "a caption"}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(Path("unused-template.pptx"), plan)


def test_fill_template_bool_layout_ref_raises(template_path: Path):
    """``bool`` is an ``int`` subclass in Python -- ``true``/``false`` must
    still be rejected as a layout reference, not silently treated as 1/0."""
    plan = {"slides": [{"layout": True, "placeholders": {}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_missing_layout_key_raises(template_path: Path):
    plan = {"slides": [{"placeholders": {}}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_non_dict_content_plan_raises(template_path: Path):
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, ["not", "a", "dict"])


def test_fill_template_missing_slides_key_raises(template_path: Path):
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, {})


def test_fill_template_produces_one_slide_per_entry_with_real_text(
    template_path: Path,
):
    plan = {
        "slides": [
            {"layout": 0, "placeholders": {"0": "Title A", "1": "Subtitle A"}},
            {
                "layout": 1,
                "placeholders": {"0": "Title B", "1": ["line one", "line two"]},
            },
        ]
    }
    prs = pptx_pipeline.fill_template(template_path, plan)
    assert len(prs.slides) == 2
    slide0, slide1 = prs.slides
    assert slide0.placeholders[0].text_frame.text == "Title A"
    assert slide0.placeholders[1].text_frame.text == "Subtitle A"
    assert slide1.placeholders[0].text_frame.text == "Title B"
    assert slide1.placeholders[1].text_frame.text == "line one\nline two"


# === run_spec / run_fill (file-path orchestration) ===========================


def test_run_spec_returns_spec_and_optionally_writes_json(template_path: Path, tmp_path: Path):
    spec = pptx_pipeline.run_spec(template_path)
    assert len(spec.layouts) == 11

    out_path = tmp_path / "spec.json"
    spec_written = pptx_pipeline.run_spec(template_path, out_path=out_path)
    assert out_path.is_file()
    assert json.loads(out_path.read_text(encoding="utf-8")) == pptx_pipeline.spec_to_dict(spec_written)


def test_run_fill_missing_template_raises_and_writes_no_output(tmp_path: Path):
    plan_path = _write_content_plan(tmp_path, {"slides": []})
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.PptxTemplateError):
        pptx_pipeline.run_fill(tmp_path / "no-such-template.pptx", plan_path, out_path)
    assert not out_path.exists()


def test_run_fill_unreadable_content_plan_path_raises(template_path: Path, tmp_path: Path):
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, tmp_path / "no-such-plan.json", out_path)
    assert not out_path.exists()


def test_run_fill_malformed_json_content_plan_raises(template_path: Path, tmp_path: Path):
    plan_path = tmp_path / "content_plan.json"
    plan_path.write_text("{not valid json", encoding="utf-8")
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert not out_path.exists()


def test_run_fill_invalid_utf8_content_plan_raises_and_writes_no_output(template_path: Path, tmp_path: Path):
    """``UnicodeDecodeError`` is a ``ValueError`` subclass, not an
    ``OSError`` subclass -- a non-UTF-8 content plan must still raise the
    module's own ``InvalidContentPlanError``, not the raw stdlib
    exception."""
    plan_path = tmp_path / "content_plan.json"
    plan_path.write_bytes(b"\xff\xfe not valid utf-8")
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert not out_path.exists()


def test_run_fill_unknown_placeholder_idx_writes_no_output_file(template_path: Path, tmp_path: Path):
    plan_path = _write_content_plan(tmp_path, {"slides": [{"layout": 0, "placeholders": {"99": "nope"}}]})
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert not out_path.exists()


def test_run_fill_writes_a_real_pptx_file(template_path: Path, tmp_path: Path):
    plan_path = _write_content_plan(tmp_path, {"slides": [{"layout": 0, "placeholders": {"0": "Hi", "1": "There"}}]})
    out_path = tmp_path / "out.pptx"
    pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert out_path.is_file()
    prs = Presentation(str(out_path))
    assert len(prs.slides) == 1


# === Story 15.2: autofit engine ===============================================


def test_wrap_words_fits_on_one_line_when_width_is_generous():
    assert pptx_pipeline._wrap_words("hello world", 12, 1000.0) == ["hello world"]


def test_wrap_words_wraps_at_a_word_boundary_when_width_is_tight():
    single_word_width_pt = pptx_pipeline._load_font(12).getlength("hello")
    lines = pptx_pipeline._wrap_words("hello world", 12, single_word_width_pt + 1)
    assert lines == ["hello", "world"]


def test_wrap_words_places_an_overlong_single_word_alone_rather_than_splitting():
    lines = pptx_pipeline._wrap_words("supercalifragilistic", 12, 1.0)
    assert lines == ["supercalifragilistic"]


def test_wrap_words_empty_text_returns_one_empty_line():
    assert pptx_pipeline._wrap_words("", 12, 100.0) == [""]


def test_fits_at_size_true_when_line_is_narrower_than_width():
    width_pt = pptx_pipeline._load_font(12).getlength("short") + 5
    assert pptx_pipeline._fits_at_size("short", 12, width_pt) is True


def test_fits_at_size_false_when_line_is_wider_than_width():
    text = "a much longer line of text than the width allows"
    width_pt = pptx_pipeline._load_font(12).getlength(text) - 5
    assert pptx_pipeline._fits_at_size(text, 12, width_pt) is False


def test_rebalance_orphan_merges_a_single_word_orphan_when_the_merge_fits():
    width_pt = pptx_pipeline._load_font(12).getlength("bar baz") + 1
    assert pptx_pipeline._rebalance_orphan(["foo bar", "baz"], 12, width_pt) == [
        "foo",
        "bar baz",
    ]


def test_rebalance_orphan_leaves_a_multi_word_last_line_untouched():
    lines = ["foo", "bar baz"]
    assert pptx_pipeline._rebalance_orphan(lines, 12, 1000.0) == lines


def test_rebalance_orphan_leaves_a_single_line_untouched():
    assert pptx_pipeline._rebalance_orphan(["only one line"], 12, 1000.0) == ["only one line"]


def test_rebalance_orphan_no_op_when_prior_line_has_only_one_word():
    assert pptx_pipeline._rebalance_orphan(["foo", "bar"], 12, 1000.0) == ["foo", "bar"]


def test_rebalance_orphan_no_op_when_the_merge_does_not_fit_width():
    lines = ["foo bar", "baz"]
    assert pptx_pipeline._rebalance_orphan(lines, 12, 1.0) == lines


def test_fit_text_uses_max_pt_when_text_fits_comfortably():
    fitted = pptx_pipeline.fit_text("hi", width_emu=5_000_000, height_emu=5_000_000, max_pt=24, min_pt=6)
    assert fitted.font_size_pt == 24
    assert fitted.lines == ("hi",)


def test_fit_text_shrinks_below_max_pt_to_fit_a_tight_box():
    long_text = "word " * 20
    fitted = pptx_pipeline.fit_text(long_text, width_emu=900_000, height_emu=900_000, max_pt=40, min_pt=6)
    height_budget_pt = (900_000 / 12700) * 0.9
    assert fitted.font_size_pt < 40
    assert len(fitted.lines) * fitted.font_size_pt * 1.2 <= height_budget_pt


def test_fit_text_extreme_overflow_renders_at_min_pt_without_raising():
    """I/O matrix's "Extreme overflow" row: text so long even the
    minimum font size still overflows the box -- renders at the minimum
    size anyway, no error, no crash."""
    huge_text = "word " * 500
    fitted = pptx_pipeline.fit_text(huge_text, width_emu=50_000, height_emu=50_000, max_pt=20, min_pt=6)
    assert fitted.font_size_pt == 6
    assert len(fitted.lines) > 0


def test_fit_text_shrinks_an_unbreakable_token_to_fit_the_width():
    """A single word wider than the box is deliberately left alone on its
    own line, so it clears any height budget while still running past the
    box's right edge. The search must reject those sizes on the width
    axis too -- otherwise a long unbroken metric value, URL, or hash
    "fits" at nearly full size and overflows horizontally."""
    # A box tall enough that height alone would happily accept max_pt,
    # but only wide enough for the token at a much smaller size.
    width_emu = 700_000
    fitted = pptx_pipeline.fit_text("1,247,392", width_emu=width_emu, height_emu=3_000_000, max_pt=54, min_pt=6)
    width_budget_pt = (width_emu / 12700) * pptx_pipeline._FIT_SAFETY
    assert len(fitted.lines) == 1
    assert fitted.font_size_pt < 54
    assert pptx_pipeline._load_font(fitted.font_size_pt).getlength(fitted.lines[0]) <= width_budget_pt


def test_fit_text_falls_back_to_min_pt_when_a_token_cannot_fit_any_width():
    """The width axis must not break the "Extreme overflow" contract: a
    token too wide even at `min_pt` still renders best-effort at the
    floor rather than raising or looping."""
    fitted = pptx_pipeline.fit_text("1,247,392", width_emu=300_000, height_emu=300_000, max_pt=54, min_pt=6)
    assert fitted.font_size_pt == 6
    assert fitted.lines == ("1,247,392",)


def test_fit_text_applies_the_safety_margin_to_width_as_well_as_height():
    """The 10% margin absorbs both the metrics gap against whatever font
    PowerPoint substitutes and the extra advance of the bold faces used
    for card titles, metric values, and table headers -- which are
    measured with Pillow's regular face. It must therefore hold on the
    width axis, not only the height axis."""
    width_emu = 1_200_000
    fitted = pptx_pipeline.fit_text("measured to fit", width_emu=width_emu, height_emu=5_000_000, max_pt=40, min_pt=6)
    width_budget_pt = (width_emu / 12700) * pptx_pipeline._FIT_SAFETY
    for line in fitted.lines:
        assert pptx_pipeline._load_font(fitted.font_size_pt).getlength(line) <= width_budget_pt


def test_lines_height_emu_matches_the_line_height_formula():
    fitted = pptx_pipeline.FittedText(font_size_pt=12, lines=("a", "b", "c"))
    assert pptx_pipeline._lines_height_emu(fitted) == round(3 * 12 * 1.2 * 12700)


# === Story 15.2: shape functions (add_card / add_metric_box / add_table /
# add_section_label) ============================================================


def _blank_slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def test_add_card_creates_a_real_autoshape_with_theme_colors_and_real_text():
    _, slide = _blank_slide()
    shape = pptx_pipeline.add_card(slide, 0, 0, 2_000_000, 1_500_000, "A Title", "A body sentence.")
    assert shape.has_text_frame
    assert len(shape.text_frame.paragraphs) == 2
    assert "A Title" in shape.text_frame.text
    assert "A body sentence." in shape.text_frame.text
    xml = shape._element.xml
    assert "schemeClr" in xml
    assert "srgbClr" not in xml

    title_run = shape.text_frame.paragraphs[0].runs[0]
    body_run = shape.text_frame.paragraphs[1].runs[0]
    assert title_run.font.bold is True
    assert body_run.font.bold is False

    # `_reset_margins` must actually zero all four margins -- `fit_text`'s
    # EMU width/height budget assumes the shape's box IS the text area.
    text_frame = shape.text_frame
    assert text_frame.margin_left == 0
    assert text_frame.margin_right == 0
    assert text_frame.margin_top == 0
    assert text_frame.margin_bottom == 0


def test_add_metric_box_renders_a_large_value_line_above_a_small_label_line():
    _, slide = _blank_slide()
    shape = pptx_pipeline.add_metric_box(slide, 0, 0, 1_500_000, 900_000, "42%", "of findings resolved")
    value_paragraph, label_paragraph = shape.text_frame.paragraphs
    value_run = value_paragraph.runs[0]
    label_run = label_paragraph.runs[0]
    assert value_run.text == "42%"
    assert "of findings resolved" in label_paragraph.text
    assert value_run.font.size.pt > label_run.font.size.pt

    assert value_run.font.bold is True
    assert value_run.font.color.theme_color == pptx_pipeline.MSO_THEME_COLOR.ACCENT_1
    assert label_run.font.bold is False
    assert label_run.font.color.theme_color == pptx_pipeline.MSO_THEME_COLOR.TEXT_1

    text_frame = shape.text_frame
    assert text_frame.margin_left == 0
    assert text_frame.margin_right == 0
    assert text_frame.margin_top == 0
    assert text_frame.margin_bottom == 0


def test_add_table_gives_every_cell_one_font_size_and_a_distinct_header_row():
    _, slide = _blank_slide()
    rows = (("Axis", "Signal"), ("Hygiene", "deptry"), ("Security", "osv-scanner"))
    graphic_frame = pptx_pipeline.add_table(slide, 0, 0, 4_000_000, 1_500_000, rows)
    table = graphic_frame.table
    sizes = {
        table.cell(r, c).text_frame.paragraphs[0].runs[0].font.size.pt
        for r in range(len(rows))
        for c in range(len(rows[0]))
    }
    assert len(sizes) == 1  # one global size shared by every cell

    # The header's readability is a *contrast pair*: light theme text on
    # an ACCENT_1 fill. A substring check for "schemeClr" cannot see it --
    # the run's own colour already puts that string in the cell XML -- so
    # assert both halves of the pair by identity.
    header_cell = table.cell(0, 0)
    header_run = header_cell.text_frame.paragraphs[0].runs[0]
    assert header_run.font.bold is True
    assert header_run.font.color.theme_color == pptx_pipeline.MSO_THEME_COLOR.BACKGROUND_1
    assert header_cell.fill.fore_color.theme_color == pptx_pipeline.MSO_THEME_COLOR.ACCENT_1

    body_cell = table.cell(1, 0)
    body_run = body_cell.text_frame.paragraphs[0].runs[0]
    assert body_run.font.bold is False
    assert body_run.font.color.theme_color == pptx_pipeline.MSO_THEME_COLOR.TEXT_1

    # `_reset_margins` must actually zero all four margins on a cell's
    # text frame too, same invariant as the other shape functions.
    body_text_frame = body_cell.text_frame
    assert body_text_frame.margin_left == 0
    assert body_text_frame.margin_right == 0
    assert body_text_frame.margin_top == 0
    assert body_text_frame.margin_bottom == 0


def test_every_written_paragraph_pins_spacing_and_alignment():
    """`fit_text`'s height budget is only true on the rendered slide
    while inter-paragraph spacing stays zero and the line-height stays
    single -- `_write_fitted_lines` pins all three rather than inheriting
    them from the theme, so a template swap cannot silently re-inflate a
    measured-to-fit shape. It pins `alignment` for the same reason:
    python-pptx's autoshape template ships a CENTER first paragraph while
    an added paragraph inherits the theme default, so an unpinned card
    would render a centred title over a left-aligned body."""
    _, slide = _blank_slide()
    card = pptx_pipeline.add_card(slide, 0, 0, 2_000_000, 1_500_000, "A Title", "A body sentence.")
    metric = pptx_pipeline.add_metric_box(slide, 0, 0, 1_500_000, 900_000, "42%", "of findings resolved")
    label = pptx_pipeline.add_section_label(slide, 0, 0, 3_000_000, 400_000, "Act I")
    table = pptx_pipeline.add_table(
        slide, 0, 0, 4_000_000, 1_500_000, (("Axis", "Signal"), ("Hygiene", "deptry"))
    ).table

    paragraphs = [
        *card.text_frame.paragraphs,
        *metric.text_frame.paragraphs,
        *label.text_frame.paragraphs,
        *(table.cell(r, c).text_frame.paragraphs[0] for r in range(2) for c in range(2)),
    ]
    for paragraph in paragraphs:
        assert paragraph.space_before == 0
        assert paragraph.space_after == 0
        assert paragraph.line_spacing == 1.0
        assert paragraph.alignment is not None

    # Within one shape, every paragraph agrees on alignment.
    assert len({p.alignment for p in card.text_frame.paragraphs}) == 1
    assert len({p.alignment for p in metric.text_frame.paragraphs}) == 1


def test_add_card_title_never_renders_smaller_than_its_body():
    """A long title shrinks; the body must not then out-size it and
    invert the card's visual hierarchy. The body's search ceiling is
    clamped to whatever size the title settled on."""
    _, slide = _blank_slide()
    shape = pptx_pipeline.add_card(
        slide,
        0,
        0,
        2_000_000,
        1_500_000,
        "A Title Here That Is Long Enough To Wrap Nicely",
        "short body",
    )
    title_run = shape.text_frame.paragraphs[0].runs[0]
    body_run = shape.text_frame.paragraphs[1].runs[0]
    assert title_run.font.size.pt >= body_run.font.size.pt


def test_add_metric_box_label_never_renders_larger_than_its_value():
    """The mirror of the card's clamp: a long value forced small by a
    narrow box must not end up smaller than the short label annotating
    it. Unclamped, this fixture renders a 9pt value under a 14pt label."""
    _, slide = _blank_slide()
    shape = pptx_pipeline.add_metric_box(slide, 0, 0, 600_000, 900_000, "1,247,392", "pkgs")
    value_run = shape.text_frame.paragraphs[0].runs[0]
    label_run = shape.text_frame.paragraphs[1].runs[0]
    assert value_run.font.size.pt >= label_run.font.size.pt


def test_add_metric_box_label_fits_the_height_left_by_its_value():
    """The adaptive split must leave the label real room: value block +
    label block together stay within the box (the direct counterpart of
    the card's Warden acceptance check)."""
    _, slide = _blank_slide()
    height = 900_000
    shape = pptx_pipeline.add_metric_box(slide, 0, 0, 1_500_000, height, "42%", "of findings resolved this quarter")
    total_emu = 0
    for paragraph in shape.text_frame.paragraphs:
        line_count = paragraph.text.count("\x0b") + 1
        size_pt = paragraph.runs[0].font.size.pt
        total_emu += round(line_count * size_pt * 1.2 * 12700)
    assert total_emu <= height


def test_add_table_shrinks_to_its_own_per_row_height_budget():
    """`add_table` budgets each cell against `height / num_rows`, not the
    whole table height. A many-rowed table of long cells must therefore
    come back small enough that each cell's own text block fits one row."""
    _, slide = _blank_slide()
    width, height = 8_000_000, 2_000_000
    rows = tuple((f"Row {index} axis label", f"a reasonably long signal description {index}") for index in range(5))
    graphic_frame = pptx_pipeline.add_table(slide, 0, 0, width, height, rows)
    table = graphic_frame.table

    row_budget_pt = (height / len(rows)) / 12700
    width_budget_pt = ((width / len(rows[0])) / 12700) * pptx_pipeline._FIT_SAFETY
    for row_index in range(len(rows)):
        for col_index in range(len(rows[0])):
            paragraph = table.cell(row_index, col_index).text_frame.paragraphs[0]
            size_pt = paragraph.runs[0].font.size.pt
            line_count = paragraph.text.count("\x0b") + 1
            assert line_count * size_pt * 1.2 <= row_budget_pt
            for run in paragraph.runs:
                assert pptx_pipeline._load_font(int(size_pt)).getlength(run.text) <= width_budget_pt


def test_add_table_single_row_has_no_header():
    _, slide = _blank_slide()
    graphic_frame = pptx_pipeline.add_table(slide, 0, 0, 2_000_000, 500_000, (("only", "row"),))
    cell = graphic_frame.table.cell(0, 0)
    assert cell.text_frame.paragraphs[0].runs[0].font.bold is False


def test_table_font_size_falls_back_to_min_pt_when_nothing_fits():
    """When no candidate size in `_table_font_size`'s descending search
    range lets a cell's text fit its own budget -- even at `_MIN_PT` --
    the function still falls through to the shared minimum rather than
    raising or looping forever."""
    rows = (("This is a very long cell of text that will not fit", "b"),)
    assert pptx_pipeline._table_font_size(rows, cell_width_emu=50_000, cell_height_emu=50_000) == pptx_pipeline._MIN_PT


def test_add_section_label_renders_a_single_theme_colored_run():
    _, slide = _blank_slide()
    textbox = pptx_pipeline.add_section_label(slide, 0, 0, 3_000_000, 400_000, "Act I")
    assert textbox.text_frame.text == "Act I"
    run = textbox.text_frame.paragraphs[0].runs[0]
    assert run.font.color.theme_color == pptx_pipeline.MSO_THEME_COLOR.ACCENT_1
    assert run.font.bold is True

    text_frame = textbox.text_frame
    assert text_frame.margin_left == 0
    assert text_frame.margin_right == 0
    assert text_frame.margin_top == 0
    assert text_frame.margin_bottom == 0


# === Story 15.2: content_plan.json "shapes" list (fill_template / run_fill) ===


def _shape_plan(shape: dict, *, layout: int = 6) -> dict:
    return {"slides": [{"layout": layout, "placeholders": {}, "shapes": [shape]}]}


def test_fill_template_card_shape_produces_real_text_runs(template_path: Path):
    plan = _shape_plan(
        {
            "type": "card",
            "left": 0,
            "top": 0,
            "width": 2_000_000,
            "height": 1_500_000,
            "title": "A Title",
            "body": "A body sentence.",
        }
    )
    prs = pptx_pipeline.fill_template(template_path, plan)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    card = next(s for s in slide.shapes if s.has_text_frame and "A Title" in s.text_frame.text)
    assert "A body sentence." in card.text_frame.text


@pytest.mark.parametrize(
    "shape",
    [
        {
            "type": "metric_box",
            "left": 0,
            "top": 0,
            "width": 1_500_000,
            "height": 900_000,
            "value": "42%",
            "label": "of findings resolved",
        },
        {
            "type": "table",
            "left": 0,
            "top": 0,
            "width": 3_000_000,
            "height": 900_000,
            "rows": [["Axis", "Signal"], ["Hygiene", "deptry"]],
        },
        {
            "type": "section_label",
            "left": 0,
            "top": 0,
            "width": 2_000_000,
            "height": 400_000,
            "text": "Act I",
        },
    ],
)
def test_fill_template_every_shape_type_produces_a_slide(template_path: Path, shape):
    prs = pptx_pipeline.fill_template(template_path, _shape_plan(shape))
    assert len(prs.slides) == 1


def test_run_fill_shapes_produce_real_runs_no_pic_and_theme_colors_only(template_path: Path, tmp_path: Path):
    """Acceptance: the output .pptx's slide XML contains real <a:t> runs
    for every shape's text (no <p:pic> covering them), and every shape's
    fill/text color is a theme color reference (<a:schemeClr>), never a
    hardcoded hex value."""
    plan = {
        "slides": [
            {
                "layout": 6,
                "placeholders": {},
                "shapes": [
                    {
                        "type": "card",
                        "left": 0,
                        "top": 0,
                        "width": 2_000_000,
                        "height": 1_500_000,
                        "title": "Card Title",
                        "body": "Card body text.",
                    },
                    {
                        "type": "metric_box",
                        "left": 0,
                        "top": 1_600_000,
                        "width": 1_500_000,
                        "height": 900_000,
                        "value": "42%",
                        "label": "of findings resolved",
                    },
                    {
                        "type": "table",
                        "left": 0,
                        "top": 2_600_000,
                        "width": 3_000_000,
                        "height": 900_000,
                        "rows": [["Axis", "Signal"], ["Hygiene", "deptry"]],
                    },
                    {
                        "type": "section_label",
                        "left": 0,
                        "top": 3_600_000,
                        "width": 2_000_000,
                        "height": 400_000,
                        "text": "Act I",
                    },
                ],
            }
        ]
    }
    plan_path = _write_content_plan(tmp_path, plan)
    out_path = tmp_path / "out.pptx"

    pptx_pipeline.run_fill(template_path, plan_path, out_path)

    xml = _slide_xml(out_path, 1)
    assert "<p:pic" not in xml
    for text in (
        "Card Title",
        "Card body text.",
        "42%",
        "of findings resolved",
        "Axis",
        "Act I",
    ):
        assert text in xml
    assert "<a:schemeClr" in xml
    assert "srgbClr" not in xml


def test_fill_template_shapes_not_a_list_raises(template_path: Path):
    plan = {"slides": [{"layout": 6, "placeholders": {}, "shapes": "not-a-list"}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_unknown_shape_type_raises_and_writes_no_output(template_path: Path, tmp_path: Path):
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan({"type": "chart", "left": 0, "top": 0, "width": 1, "height": 1}),
    )
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert not out_path.exists()


def test_fill_template_card_shape_missing_body_raises(template_path: Path):
    plan = _shape_plan(
        {
            "type": "card",
            "left": 0,
            "top": 0,
            "width": 1_000_000,
            "height": 1_000_000,
            "title": "x",
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


@pytest.mark.parametrize(
    "geometry",
    [
        {"left": 0, "top": 0, "width": 0, "height": 1_000_000},
        {"left": 0, "top": 0, "width": 1_000_000, "height": -1},
        {"left": -1, "top": 0, "width": 1_000_000, "height": 1_000_000},
    ],
)
def test_fill_template_non_positive_shape_geometry_raises(template_path: Path, geometry):
    plan = _shape_plan({"type": "section_label", "text": "x", **geometry})
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_shape_geometry_bool_value_raises(template_path: Path):
    """``bool`` is an ``int`` subclass in Python -- a shape's geometry
    field must still reject ``true``/``false``, not silently treat it as
    1/0 (the same discipline `_resolve_layout`'s layout-reference check
    already applies, per ``_resolve_shape_geometry``'s own docstring)."""
    plan = _shape_plan(
        {
            "type": "section_label",
            "text": "x",
            "left": 0,
            "top": 0,
            "width": True,
            "height": 1_000_000,
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


@pytest.mark.parametrize(
    "rows",
    [
        "not-a-list",
        [],
        [["a", "b"], ["c"]],  # non-rectangular: differing column counts
        [["a", 1]],  # non-string cell
    ],
)
def test_fill_template_table_rows_validation_raises(template_path: Path, rows):
    plan = _shape_plan(
        {
            "type": "table",
            "left": 0,
            "top": 0,
            "width": 2_000_000,
            "height": 1_000_000,
            "rows": rows,
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_run_fill_invalid_table_rows_writes_no_output_file(template_path: Path, tmp_path: Path):
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan(
            {
                "type": "table",
                "left": 0,
                "top": 0,
                "width": 2_000_000,
                "height": 1_000_000,
                "rows": [["a", "b"], ["c"]],  # non-rectangular
            }
        ),
    )
    out_path = tmp_path / "out.pptx"
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.run_fill(template_path, plan_path, out_path)
    assert not out_path.exists()


def test_fill_template_metric_box_missing_value_raises(template_path: Path):
    plan = _shape_plan(
        {
            "type": "metric_box",
            "left": 0,
            "top": 0,
            "width": 1_000_000,
            "height": 1_000_000,
            "label": "of findings resolved",
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_metric_box_missing_label_raises(template_path: Path):
    plan = _shape_plan(
        {
            "type": "metric_box",
            "left": 0,
            "top": 0,
            "width": 1_000_000,
            "height": 1_000_000,
            "value": "42%",
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_section_label_missing_text_raises(template_path: Path):
    plan = _shape_plan(
        {
            "type": "section_label",
            "left": 0,
            "top": 0,
            "width": 1_000_000,
            "height": 1_000_000,
        }
    )
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_shape_missing_type_key_raises_a_distinct_message(
    template_path: Path,
):
    """A shape entry missing the ``"type"`` key entirely must raise with
    a message distinct from an entry carrying an unrecognized ``"type"``
    value (both would otherwise read identically: "unknown shape type
    None")."""
    plan = _shape_plan({"left": 0, "top": 0, "width": 1_000_000, "height": 1_000_000})
    with pytest.raises(errors.InvalidContentPlanError, match="missing 'type'"):
        pptx_pipeline.fill_template(template_path, plan)


@pytest.mark.parametrize("shape", ["not-a-mapping", 42])
def test_fill_template_non_mapping_shape_entry_raises(template_path: Path, shape):
    plan = {"slides": [{"layout": 6, "placeholders": {}, "shapes": [shape]}]}
    with pytest.raises(errors.InvalidContentPlanError):
        pptx_pipeline.fill_template(template_path, plan)


def test_fill_template_no_shapes_key_behaves_exactly_as_story_15_1(template_path: Path):
    """No `"shapes"` key -- an existing Story 15.1-shaped content_plan
    entry (placeholders only) -- must behave exactly as before this
    story."""
    plan = {"slides": [{"layout": 0, "placeholders": {"0": "Hi", "1": "There"}}]}
    prs = pptx_pipeline.fill_template(template_path, plan)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert slide.placeholders[0].text_frame.text == "Hi"
    assert slide.placeholders[1].text_frame.text == "There"


def test_fill_template_placeholders_and_shapes_coexist_on_the_same_slide(
    template_path: Path,
):
    """The `"shapes"` key is additive to a slide entry's existing
    `"placeholders"` mapping -- a single slide entry setting BOTH a real
    (non-empty) `"placeholders"` mapping AND a real `"shapes"` list must
    materialize both with no interference."""
    plan = {
        "slides": [
            {
                "layout": 0,
                "placeholders": {"0": "A Real Title", "1": "A Real Subtitle"},
                "shapes": [
                    {
                        "type": "card",
                        "left": 0,
                        "top": 0,
                        "width": 2_000_000,
                        "height": 1_500_000,
                        "title": "Card Title",
                        "body": "Card body text.",
                    }
                ],
            }
        ]
    }
    prs = pptx_pipeline.fill_template(template_path, plan)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert slide.placeholders[0].text_frame.text == "A Real Title"
    assert slide.placeholders[1].text_frame.text == "A Real Subtitle"
    card = next(s for s in slide.shapes if s.has_text_frame and "Card Title" in s.text_frame.text)
    assert "Card body text." in card.text_frame.text


# === Story 15.2: the Warden-appendix acceptance case ===========================

_WARDEN_APPENDIX_PERSONAS = [
    (
        "CISO — provable risk posture",
        "exploitability-prioritized findings (KEV/EPSS) and audit-grade, reproducible evidence.",
    ),
    (
        "Chief Dev Experience — a gate devs trust",
        "runs locally first, one command, no false alarms — so it doesn't cry wolf.",
    ),
    (
        "CIO — standardized & offline",
        "one gate fleet-wide, deterministic and air-gap-ready for regulated estates.",
    ),
    (
        "Chief Data & Analytics — the ML footprint, covered",
        "conda / conda-forge coverage — the scientific / ML library estate that stock PyPI-only tools never parse.",
    ),
]
"""The 4 real Appendix persona strings, transcribed verbatim (minus
markdown `**` bold markers, split at the `:**` boundary into
title/body) from
presentations/pyforge-warden/src/marp/pyforge-warden-deck-2026-07-15.md:413-419
-- the only `## Appendix` slide in any station deck in this repo."""


def test_add_card_fits_the_warden_appendix_personas_within_card_height():
    """Acceptance: given the 4 real Warden-deck Appendix persona strings
    packed into four 2.15in x 3.0in card shapes, when add_card computes
    their autofit, then every card's rendered text block height,
    re-measured independently of add_card's internal computation
    (paragraph line-break count x font size x 1.2 -- the same formula
    fit_text's own search budgets against), stays within the card's
    height."""
    card_width_emu = round(2.15 * 914400)
    card_height_emu = round(3.0 * 914400)
    card_height_pt = card_height_emu / 12700

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    for title, body in _WARDEN_APPENDIX_PERSONAS:
        shape = pptx_pipeline.add_card(slide, 0, 0, card_width_emu, card_height_emu, title, body)
        total_height_pt = 0.0
        for paragraph in shape.text_frame.paragraphs:
            font_size_pt = paragraph.runs[0].font.size.pt
            line_count = paragraph.text.count("\x0b") + 1
            total_height_pt += line_count * font_size_pt * 1.2
        assert total_height_pt <= card_height_pt


# === CLI wiring ================================================================


def test_deck_pptx_spec_help_exits_zero():
    assert cli.main(["deck", "pptx-spec", "--help"]) == 0


def test_deck_pptx_fill_help_exits_zero():
    assert cli.main(["deck", "pptx-fill", "--help"]) == 0


def test_deck_pptx_spec_prints_json_to_stdout_by_default(capsys):
    exit_code = cli.main(["deck", "pptx-spec"])
    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert len(payload["layouts"]) == 11


def test_deck_pptx_spec_writes_to_out_path(tmp_path: Path, capsys):
    out_path = tmp_path / "spec.json"
    exit_code = cli.main(["deck", "pptx-spec", "-o", str(out_path)])
    assert exit_code == 0
    assert out_path.is_file()
    assert "wrote" in capsys.readouterr().out


def test_deck_pptx_spec_missing_template_exits_1(tmp_path: Path, capsys):
    exit_code = cli.main(["deck", "pptx-spec", "--template", str(tmp_path / "nope.pptx")])
    assert exit_code == 1
    assert "PptxTemplateError" in capsys.readouterr().err


def test_deck_pptx_fill_writes_pptx(tmp_path: Path, capsys):
    plan_path = _write_content_plan(tmp_path, {"slides": [{"layout": 0, "placeholders": {"0": "Hi", "1": "There"}}]})
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 0
    assert out_path.is_file()
    assert "wrote" in capsys.readouterr().out


def test_deck_pptx_fill_unknown_placeholder_idx_exits_1_and_writes_nothing(tmp_path: Path, capsys):
    plan_path = _write_content_plan(tmp_path, {"slides": [{"layout": 0, "placeholders": {"99": "nope"}}]})
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 1
    assert not out_path.exists()
    assert "InvalidContentPlanError" in capsys.readouterr().err


def test_deck_pptx_fill_unknown_shape_type_exits_1_and_writes_nothing(tmp_path: Path, capsys):
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan({"type": "chart", "left": 0, "top": 0, "width": 1, "height": 1}),
    )
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 1
    assert not out_path.exists()
    assert "InvalidContentPlanError" in capsys.readouterr().err


def test_deck_pptx_fill_missing_required_field_exits_1_and_writes_nothing(tmp_path: Path, capsys):
    """I/O matrix's "Missing required field" row, proven all the way
    through the CLI (mirroring the unknown-shape-type CLI test above) --
    not just at the in-memory `fill_template` level."""
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan(
            {"type": "card", "left": 0, "top": 0, "width": 1_000_000, "height": 1_000_000, "title": "x"}  # no "body"
        ),
    )
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 1
    assert not out_path.exists()
    assert "InvalidContentPlanError" in capsys.readouterr().err


def test_deck_pptx_fill_non_positive_shape_geometry_exits_1_and_writes_nothing(tmp_path: Path, capsys):
    """I/O matrix's "Non-positive geometry" row, proven all the way
    through the CLI (mirroring the unknown-shape-type CLI test above) --
    not just at the in-memory `fill_template` level."""
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan({"type": "section_label", "text": "x", "left": 0, "top": 0, "width": 0, "height": 1_000_000}),
    )
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 1
    assert not out_path.exists()
    assert "InvalidContentPlanError" in capsys.readouterr().err


def test_deck_pptx_fill_cli_handles_a_shapes_bearing_plan(tmp_path: Path, capsys):
    plan_path = _write_content_plan(
        tmp_path,
        _shape_plan(
            {
                "type": "section_label",
                "left": 0,
                "top": 0,
                "width": 2_000_000,
                "height": 400_000,
                "text": "Act I",
            }
        ),
    )
    out_path = tmp_path / "out.pptx"
    exit_code = cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)])
    assert exit_code == 0
    assert out_path.is_file()
    assert "wrote" in capsys.readouterr().out


def test_deck_pptx_subcommands_never_construct_an_mcp_transport(monkeypatch, tmp_path: Path):
    """Mirrors ``test_cli_deck_qa.py``'s regression: these subcommands are
    fully local/offline and must never reach Claude Design."""

    class _ExplodingTransport:
        def __init__(self, *args, **kwargs):
            raise AssertionError("pptx-spec/pptx-fill must never construct McpTransport")

    monkeypatch.setattr(cli, "McpTransport", _ExplodingTransport)

    assert cli.main(["deck", "pptx-spec"]) == 0

    plan_path = _write_content_plan(tmp_path, {"slides": [{"layout": 0, "placeholders": {"0": "Hi"}}]})
    out_path = tmp_path / "out.pptx"
    assert cli.main(["deck", "pptx-fill", str(plan_path), "-o", str(out_path)]) == 0


# === the flagship round-trip proof ============================================


def test_round_trip_three_real_slides_produce_genuinely_editable_text(template_path: Path, tmp_path: Path):
    """Given the bundled default template and a content_plan.json built
    from 3 real slides transcribed from
    presentations/pyforge-herald/src/marp/pyforge-herald-deck-2026-07-24.md
    (cover, "Invisible engineering is failed engineering", "The
    Design-Code bridge -- realized" bullets), when run_fill is invoked,
    the output .pptx contains those exact strings as real <a:t> text runs
    (no <p:pic> covering the filled placeholders) AND the same text
    survives an independent `soffice --headless --convert-to pptx`
    round-trip."""
    plan_path = _write_content_plan(tmp_path, _THREE_REAL_SLIDES_PLAN)
    out_path = tmp_path / "herald-deck.pptx"

    pptx_pipeline.run_fill(template_path, plan_path, out_path)

    expected_strings = [
        "capture the dream. proclaim the release.",
        "Invisible engineering is failed engineering",
        "The Design–Code bridge — realized",
        "seed a Design project from the repo",
        "design visually in Claude Design",
        "pull the prototype back",
        "etags catch conflicts",
    ]

    # (1) direct OOXML proof: real <a:t> runs, no <p:pic> anywhere in any
    # of the 3 slides (Design Notes' "Round-trip proven" test proxy, step 1
    # -- catches the exact old Marp-pipeline defect directly).
    full_xml = ""
    for slide_number in (1, 2, 3):
        xml = _slide_xml(out_path, slide_number)
        assert "<p:pic" not in xml
        assert "<a:t>" in xml
        full_xml += xml
    for text in expected_strings:
        assert text in full_xml

    # (2) an independent, third-party OOXML-consuming application must
    # genuinely interpret and re-emit the same text (Design Notes, step 2,
    # the "surface-anchored" proof) -- skipped gracefully (not a story
    # failure) only if soffice is unexpectedly absent from the test
    # environment.
    soffice = shutil.which("soffice")
    if soffice is None:
        pytest.skip(
            "soffice not found on PATH -- skipping the independent "
            "round-trip proof (Design Notes: not this story's failure)"
        )

    round_trip_dir = tmp_path / "round-trip"
    round_trip_dir.mkdir()
    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pptx",
            "--outdir",
            str(round_trip_dir),
            str(out_path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert result.returncode == 0, result.stderr

    round_tripped_path = round_trip_dir / out_path.name
    assert round_tripped_path.is_file()
    round_tripped_prs = Presentation(str(round_tripped_path))
    round_tripped_text = "\n".join(
        shape.text_frame.text for slide in round_tripped_prs.slides for shape in slide.shapes if shape.has_text_frame
    )
    for text in expected_strings:
        assert text in round_tripped_text
