"""Integration: Marp model -> Node driver -> python-pptx read-back (Story 32.1)."""

from __future__ import annotations

import shutil
from datetime import datetime, timezone
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pyforge.herald import pptx_native


def _node_and_pptxgenjs_available() -> bool:
    if shutil.which("node") is None:
        return False
    prefix = __import__("os").environ.get("CONDA_PREFIX", "")
    if not prefix:
        return False
    return (Path(prefix) / "lib" / "node_modules" / "pptxgenjs-plus").is_dir()


@pytest.fixture
def _no_stamp_git(monkeypatch):
    monkeypatch.setattr("pyforge.herald.stamps.write_stamp", lambda *_a, **_k: None)


@pytest.mark.skipif(not _node_and_pptxgenjs_available(), reason="node or pptxgenjs-plus not installed")
def test_fixture_render_read_back_with_python_pptx(tmp_path: Path, monkeypatch, _no_stamp_git) -> None:
    slug = "fixture-deck"
    marp_dir = tmp_path / "presentations" / slug / "src" / "marp"
    marp_dir.mkdir(parents=True)
    marp_path = marp_dir / f"{slug}-deck-2026-09-01.md"
    marp_path.write_text(
        """---
marp: true
---

# Slide one

- alpha
- beta

<!-- notes for slide one -->

---

## Metrics

| K | V |
|---|---|
| a | 1 |

<!-- notes for table slide -->
""",
        encoding="utf-8",
    )
    fixed = datetime(2026, 9, 20, tzinfo=timezone.utc)
    result = pptx_native.export_native_pptx(
        slug,
        tmp_path,
        export_date="2026-09-20",
        now=fixed,
    )
    assert result.output_path.name == f"{slug}-deck-native-2026-09-20.pptx"
    prs = Presentation(str(result.output_path))
    assert len(prs.slides) == 2
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE
    notes = [slide.notes_slide.notes_text_frame.text for slide in prs.slides if slide.has_notes_slide]
    assert any("notes for slide one" in (text or "") for text in notes)


@pytest.mark.skipif(not _node_and_pptxgenjs_available(), reason="node or pptxgenjs-plus not installed")
def test_second_run_retires_older_native_kind(tmp_path: Path, _no_stamp_git) -> None:
    slug = "retire-deck"
    marp_dir = tmp_path / "presentations" / slug / "src" / "marp"
    marp_dir.mkdir(parents=True)
    (marp_dir / f"{slug}-deck-2026-09-01.md").write_text("# One\n", encoding="utf-8")
    older = tmp_path / "presentations" / slug / "src" / "pptx" / f"{slug}-deck-native-2026-09-01.pptx"
    older.parent.mkdir(parents=True, exist_ok=True)
    older.write_bytes(b"old-native")
    pptx_native.export_native_pptx(slug, tmp_path, export_date="2026-09-20")
    remaining = list(older.parent.glob(f"{slug}-deck-native-*.pptx"))
    assert len(remaining) == 1
    assert remaining[0].name == f"{slug}-deck-native-2026-09-20.pptx"
