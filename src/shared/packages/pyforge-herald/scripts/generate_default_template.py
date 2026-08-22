#!/usr/bin/env python
"""(Re)generate ``templates/pyforge-deck-template.pptx`` -- Story 15.1's
own-template decision (2026-08-22, spec Design Notes).

Neither named option in ``spec-pptx-deck-generation``'s open question is
executable unattended: a hand-derived "six-act visual system" ``.potx``
would require authoring custom slide *layouts*, which python-pptx cannot do
without raw-OOXML surgery (explicitly forbidden by this epic's Boundaries &
Constraints); an operator-supplied ``.potx`` requires a human who isn't
present in this run. So this story adopts python-pptx's own bundled default
Office-theme template verbatim -- ``pptx.Presentation()`` with no path
argument -- as PyForge's interim "own template": a genuine, PowerPoint-
authored ``.pptx`` with 11 real layouts and inheriting placeholders,
requiring zero hand-written OOXML.

This binary is committed to the repo, never hand-edited -- re-run this
script to regenerate it (e.g. after a python-pptx upgrade changes its own
bundled default template).

Usage:
    python scripts/generate_default_template.py
"""

from __future__ import annotations

from pptx import Presentation

from pyforge.herald.pptx_pipeline import default_template_path

_OUT_PATH = default_template_path()


def main() -> None:
    _OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    Presentation().save(str(_OUT_PATH))
    print(f"wrote {_OUT_PATH}")


if __name__ == "__main__":
    main()
