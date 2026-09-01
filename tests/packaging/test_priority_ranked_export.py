"""Story 22.1: priority.py reads identity_export_parquet and writes
identity_ranked_export.parquet for Vizro Epic 22."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pandas as pd
import pytest

pytest.importorskip(
    "openpyxl",
    reason="target scripts import openpyxl at module load; only present under -e local-recipes",
)

from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "scripts"


def _load_priority():
    script_path = SCRIPTS_DIR / "conda-forge-packaging-inventory-operations_priority.py"
    module_name = script_path.stem.replace("-", "_")
    spec = importlib.util.spec_from_file_location(module_name, script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    if str(SCRIPTS_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPTS_DIR))
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


priority = _load_priority()


def _identity_row(name: str, **extra) -> dict:
    row = {
        "Core_Python_Package_Name": name,
        "OpenTeams_Issue_URL": "",
        "conda_purl": "",
        "Conda-Forge_FeedStock_URL": "",
    }
    row.update(extra)
    return row


def _write_workbook(
    path: Path,
    *,
    identity_rows: list[dict] | None = None,
    jfrog_rows: list[dict] | None = None,
    inventory_rows: list[dict] | None = None,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ident_ws = wb.create_sheet(priority.TAB)
    ident_header = ["Core_Python_Package_Name", "OpenTeams_Issue_URL"]
    ident_ws.append(ident_header)
    for row in identity_rows or []:
        ident_ws.append([row.get(h, "") for h in ident_header])

    jfrog_ws = wb.create_sheet("CDO-ENT-JFROG")
    jfrog_header = [
        "name",
        "platform_env_count",
        "internal_app_count",
        "internal_component_count",
        "internal_lob_count",
        "artifactory_downloads",
        "artifactory_version_count",
        "risk_level",
        "vuln_status",
        "basilisk_latest_version_known_vulnerabilities_count",
    ]
    jfrog_ws.append(jfrog_header)
    for row in jfrog_rows or []:
        jfrog_ws.append([row.get(h, 0) for h in jfrog_header])

    ot_ws = wb.create_sheet("OpenTeams")
    ot_ws.append(["Title", "Priority", "URL"])

    inv_ws = wb.create_sheet("inventory-2026-08-12")
    inv_header = [
        "Core_Python_Package_Name",
        "OpenTeams_Batch",
        "OpenTeams_Cohort",
        "OpenTeams_Coverage",
        "Priority_Bucket",
    ]
    inv_ws.append(inv_header)
    for row in inventory_rows or []:
        inv_ws.append([row.get(h, "") for h in inv_header])

    wb.save(path)


def _write_identity_parquet(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_parquet(path)


def test_normal_run_writes_ranked_parquet_and_xlsx_tab(monkeypatch, tmp_path, capsys):
    xlsx_path = tmp_path / "Analysis_Dataset.xlsx"
    _write_workbook(
        xlsx_path,
        jfrog_rows=[
            {
                "name": "pkg-a",
                "platform_env_count": 2,
                "internal_app_count": 0,
                "internal_component_count": 1,
                "internal_lob_count": 0,
                "artifactory_downloads": 50,
                "artifactory_version_count": 5,
                "risk_level": "LOW",
                "vuln_status": "clean",
                "basilisk_latest_version_known_vulnerabilities_count": 0,
            }
        ],
    )

    identity_parquet = tmp_path / "identity_export.parquet"
    _write_identity_parquet(
        identity_parquet,
        [_identity_row("pkg-a", OpenTeams_Issue_URL="https://github.com/x/y/issues/1")],
    )
    ranked_export = tmp_path / "identity_ranked_export.parquet"
    canvas_path = tmp_path / "canvas.tsx"

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "prog",
            "--xlsx",
            str(xlsx_path),
            "--identity-parquet",
            str(identity_parquet),
            "--ranked-export",
            str(ranked_export),
            "--skip-inventory-sync",
            "--canvas",
            str(canvas_path),
        ],
    )

    rc = priority.main()

    assert rc == 0
    assert ranked_export.is_file()
    ranked_df = pd.read_parquet(ranked_export)
    assert list(ranked_df.columns) == priority.RANKED_EXPORT_COLUMNS
    assert ranked_df.iloc[0]["Core_Python_Package_Name"] == "pkg-a"
    assert ranked_df.iloc[0]["P"] == "P4"
    assert ranked_df.iloc[0]["Work"] == "Already tracked"
    assert ranked_df.iloc[0]["Verification_Timestamp_UTC"]

    wb = __import__("openpyxl").load_workbook(xlsx_path, data_only=True)
    ws = wb[priority.TAB]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "P"
    assert rows[1][0] == "P4"
    assert rows[1][3] == "pkg-a"
    assert "wrote" in capsys.readouterr().out


def test_missing_identity_parquet_exits_nonzero(monkeypatch, tmp_path, capsys):
    xlsx_path = tmp_path / "Analysis_Dataset.xlsx"
    _write_workbook(xlsx_path)
    missing = tmp_path / "missing.parquet"

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "prog",
            "--xlsx",
            str(xlsx_path),
            "--identity-parquet",
            str(missing),
            "--ranked-export",
            str(tmp_path / "out.parquet"),
            "--skip-inventory-sync",
        ],
    )

    with pytest.raises(SystemExit) as exc:
        priority.main()

    assert exc.value.code == 1
    err = capsys.readouterr().err
    assert str(missing) in err
    assert "pyforge-atlas-bootstrap" in err


def test_jfrog_only_package_not_ranked(monkeypatch, tmp_path):
    xlsx_path = tmp_path / "Analysis_Dataset.xlsx"
    _write_workbook(
        xlsx_path,
        jfrog_rows=[
            {
                "name": "orphan-jfrog",
                "platform_env_count": 1,
                "internal_app_count": 0,
                "internal_component_count": 0,
                "internal_lob_count": 0,
                "artifactory_downloads": 0,
                "artifactory_version_count": 0,
                "risk_level": "LOW",
                "vuln_status": "clean",
                "basilisk_latest_version_known_vulnerabilities_count": 0,
            }
        ],
    )
    identity_parquet = tmp_path / "identity_export.parquet"
    _write_identity_parquet(identity_parquet, [_identity_row("pkg-ranked")])
    ranked_export = tmp_path / "identity_ranked_export.parquet"
    canvas_path = tmp_path / "canvas.tsx"

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "prog",
            "--xlsx",
            str(xlsx_path),
            "--identity-parquet",
            str(identity_parquet),
            "--ranked-export",
            str(ranked_export),
            "--skip-inventory-sync",
            "--canvas",
            str(canvas_path),
        ],
    )

    rc = priority.main()

    assert rc == 0
    ranked_df = pd.read_parquet(ranked_export)
    assert list(ranked_df["Core_Python_Package_Name"]) == ["pkg-ranked"]
